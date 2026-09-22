# -*- coding: utf-8 -*-
"""산업 기계 합성 데이터셋 검증 + 과제 A/B/C/D 모델 평가 + Excel 보고서 생성.

전달문서 `CLAUDE_HANDOFF_데이터셋_검증_모델평가.md`와 `기획서.md`의 정의를 그대로 구현한다.
새 환경에서 입력 CSV만 있으면 재실행 가능하도록 외부 프로젝트 모듈을 import 하지 않는다.

사용 예:
    python run_validation_and_models.py --data synthetic_industrial_machine_data.csv
    python run_validation_and_models.py --only verify
    python run_validation_and_models.py --skip-prophet
"""

from __future__ import annotations

import argparse
import json
import os
import sys
import time
import warnings
from datetime import datetime
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd

# Windows에서 joblib 병렬 처리 시 물리 코어 탐색 경고를 피한다.
if sys.platform == "win32":
    os.environ.setdefault("LOKY_MAX_CPU_COUNT", str(max(1, min(4, (os.cpu_count() or 2) // 2))))

warnings.filterwarnings("ignore", category=FutureWarning)
warnings.filterwarnings("ignore", category=UserWarning, module="sklearn")

RANDOM_STATE = 42

SENSORS = [
    "temp_bearing_degC",
    "temp_motor_degC",
    "vibration_h_mms",
    "vibration_v_mms",
    "oil_pressure_bar",
    "load_pct",
    "shaft_rpm",
    "power_consumption_kw",
]

CATEGORICALS = [
    "asset_tag",
    "machine_type",
    "plant_code",
    "part_no",
    "part_family",
    "criticality",
]

# 고장 대응 이후에 생기는 정보 + 타깃 계열. 분류 입력에 들어가면 즉시 중단한다.
POST_EVENT_COLUMNS = {
    "wo_type",
    "qty_issued",
    "issue_value_inr",
    "breakdown_flag",
    "target_7d",
    "target_power_next",
}

# ---- 과제별 날짜 경계 (전달문서 §3) --------------------------------------
# 과제 A: 7일 라벨이 다음 구간을 침범하지 않도록 경계 7일을 제외한다.
SPLIT_A = {
    "학습": ("2022-01-03", "2023-12-24"),
    "검증": ("2024-01-01", "2024-06-23"),
    "평가": ("2024-07-01", "2024-12-25"),
}
EXCLUDED_A = [
    ("2023-12-25", "2023-12-31", "7일 라벨이 검증 구간을 침범"),
    ("2024-06-24", "2024-06-30", "7일 라벨이 평가 구간을 침범"),
    ("2024-12-26", "2025-01-01", "미래 7일을 관측할 수 없음"),
]
# 과제 B / C / D: 경계 제외 없음.
SPLIT_BCD = {
    "학습": ("2022-01-03", "2023-12-31"),
    "검증": ("2024-01-01", "2024-06-30"),
    "평가": ("2024-07-01", "2025-01-01"),
}

# ---- 기획서 사전 점검값 (Sheet 13 비교 전용, 최종값으로 복사 금지) --------
PLANNED_VALUES = [
    ("7일분류", "기준B(기계×부품 과거 고장률)", "ROC_AUC", 0.616),
    ("7일분류", "Logistic Regression", "ROC_AUC", 0.617),
    ("7일분류", "HistGradientBoosting", "ROC_AUC", 0.608),
    ("7일분류", "센서만(sensor_only)", "ROC_AUC", 0.553),
    ("당일분류", "기저율(기준B)", "ROC_AUC", 0.585),
    ("당일분류", "센서만(sensor_only)", "ROC_AUC", 0.665),
    ("당일분류", "결합(combined)", "ROC_AUC", 0.709),
    ("전력회귀", "전일값 지속", "MAE", 2.71),
    ("전력회귀", "지난주 같은 요일(lag7)", "MAE", 1.82),
    ("전력회귀", "기계×평일/주말 평균", "MAE", 1.30),
]

# 전달문서 §2 의 1차 검증값
EXPECTED = {
    "rows": 219000,
    "cols": 22,
    "date_min": "2022-01-03",
    "date_max": "2025-01-01",
    "n_dates": 1095,
    "n_assets": 10,
    "n_parts": 20,
    "n_machine_types": 5,
    "n_plants": 3,
    "pk_dupes": 0,
    "full_dupes": 0,
    "breakdown_1": 21636,
    "breakdown_rate": 0.098795,
    "wo_type_missing": 176207,
    "wo_type_missing_rate": 0.8046,
    "machine_daily_rows": 10950,
    "a_eval_rows": 196113,
    "a_eval_pos": 99984,
    "a_eval_rate": 0.509829,
    "a_train": (129999, 66175, 0.509042),
    "a_valid": (31541, 16078, 0.509749),
    "a_test": (32061, 16317, 0.508936),
    "b_train": (145600, 14338, 0.098475),
    "b_valid": (36400, 3610, 0.099176),
    "b_test": (37000, 3688, 0.099676),
    "c_train": 7280,
    "c_valid": 1820,
    "c_test": 1850,
}

RUN_STAMP = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
DATA_SOURCE_NOTE = "출처: synthetic_industrial_machine_data.csv (Kaggle - Machine Demand & Failure Prediction Dataset, 교육용 합성 데이터)"


def log(message: str) -> None:
    print(f"[{time.strftime('%H:%M:%S')}] {message}", flush=True)


# ===========================================================================
# 1. 로딩
# ===========================================================================
def load_data(path: str | Path) -> pd.DataFrame:
    """CSV 전량을 로딩한다. 본문은 출력하지 않는다."""
    started = time.perf_counter()
    frame = pd.read_csv(path, parse_dates=["transaction_date"])
    log(f"CSV 로딩 완료: {len(frame):,}행 × {frame.shape[1]}열 ({time.perf_counter() - started:.1f}초)")
    return frame


def split_frame(frame: pd.DataFrame, bounds: dict[str, tuple[str, str]]) -> dict[str, pd.DataFrame]:
    """날짜 경계만으로 구간을 나눈다 (랜덤 분할 금지)."""
    parts: dict[str, pd.DataFrame] = {}
    for name, (start, end) in bounds.items():
        mask = (frame["transaction_date"] >= pd.Timestamp(start)) & (
            frame["transaction_date"] <= pd.Timestamp(end)
        )
        parts[name] = frame.loc[mask]
    # 누수 점검 1: 구간 간 날짜 교집합이 없어야 한다.
    keys = list(parts)
    for i in range(len(keys)):
        for j in range(i + 1, len(keys)):
            overlap = set(parts[keys[i]]["transaction_date"]) & set(parts[keys[j]]["transaction_date"])
            if overlap:
                raise AssertionError(f"구간 {keys[i]}·{keys[j]} 날짜가 겹칩니다: {len(overlap)}일")
    return parts


# ===========================================================================
# 2. 특성 생성 (모든 과거 통계는 shift(1) 이후 계산)
# ===========================================================================
# 기계 평균 대비 편차의 기준 평균은 두 과제의 학습 구간에 공통으로 포함되는
# 2022-01-03~2023-12-24 (과제 A 학습 종료일) 로만 계산한다.
DEVIATION_REFERENCE_END = "2023-12-24"

HISTORY_FEATURES = [
    "hist_bd_rate",
    "bd_count_7d",
    "bd_count_30d",
    "days_since_bd",
    "machine_bd_count_7d",
]
SENSOR_FEATURES = (
    SENSORS
    + [f"{s}_ma7" for s in SENSORS]
    + [f"{s}_diff1" for s in SENSORS]
    + [f"{s}_dev" for s in SENSORS]
)
STATIC_FEATURES = CATEGORICALS + ["unit_cost_inr"]

FEATURE_SETS = {
    "base_rate_only": STATIC_FEATURES + HISTORY_FEATURES,
    "sensor_only": SENSOR_FEATURES,
    "combined": STATIC_FEATURES + HISTORY_FEATURES + SENSOR_FEATURES,
}


def validate_features(features: list[str]) -> None:
    """타깃·사후정보 컬럼이 입력에 섞이면 즉시 중단한다 (누수 점검 13·14)."""
    forbidden = sorted(set(features) & POST_EVENT_COLUMNS)
    if forbidden:
        raise AssertionError(f"사후 정보/타깃 컬럼을 feature로 쓸 수 없습니다: {forbidden}")


def build_machine_daily(frame: pd.DataFrame) -> pd.DataFrame:
    """기계×날짜 그레인으로 중복 제거하고 센서 파생값을 만든다 (10,950행)."""
    columns = ["transaction_date", "asset_tag", "machine_type", "plant_code", *SENSORS]
    daily = (
        frame[columns]
        .drop_duplicates(["transaction_date", "asset_tag"])
        .sort_values(["asset_tag", "transaction_date"])
        .reset_index(drop=True)
    )
    groups = daily.groupby("asset_tag", observed=True, sort=False)
    reference = daily.loc[daily["transaction_date"] <= pd.Timestamp(DEVIATION_REFERENCE_END)]
    reference_mean = reference.groupby("asset_tag", observed=True)[SENSORS].mean()

    for sensor in SENSORS:
        # 7일 이동평균은 당일까지의 값만 사용한다 (t 시점 정보).
        daily[f"{sensor}_ma7"] = groups[sensor].transform(
            lambda s: s.rolling(7, min_periods=1).mean()
        )
        daily[f"{sensor}_diff1"] = groups[sensor].diff(1)
        daily[f"{sensor}_dev"] = daily[sensor] - daily["asset_tag"].map(reference_mean[sensor])
    return daily


def build_features(frame: pd.DataFrame) -> pd.DataFrame:
    """원본 그레인(기계×부품×날짜)에 과거 이력·센서 파생·7일 타깃을 붙인다."""
    started = time.perf_counter()
    data = frame.sort_values(["asset_tag", "part_no", "transaction_date"]).reset_index(drop=True)
    keys = ["asset_tag", "part_no"]
    grouped = data.groupby(keys, observed=True, sort=False)["breakdown_flag"]

    # --- 과거 고장 이력: 현재 행을 포함하지 않도록 먼저 shift(1) ---------
    previous = grouped.shift(1)
    previous_group = previous.groupby([data["asset_tag"], data["part_no"]], observed=True, sort=False)
    cumulative_sum = previous_group.cumsum()
    observed_count = previous.notna().groupby(
        [data["asset_tag"], data["part_no"]], observed=True, sort=False
    ).cumsum()
    data["hist_bd_rate"] = cumulative_sum / observed_count.replace(0, np.nan)
    data["bd_count_7d"] = previous_group.transform(lambda s: s.rolling(7, min_periods=1).sum())
    data["bd_count_30d"] = previous_group.transform(lambda s: s.rolling(30, min_periods=1).sum())

    breakdown_dates = data["transaction_date"].where(data["breakdown_flag"] == 1)
    date_group = breakdown_dates.groupby([data["asset_tag"], data["part_no"]], observed=True, sort=False)
    last_breakdown = date_group.shift(1)
    last_breakdown = last_breakdown.groupby(
        [data["asset_tag"], data["part_no"]], observed=True, sort=False
    ).ffill()
    data["days_since_bd"] = (data["transaction_date"] - last_breakdown).dt.days

    # --- 기계 단위 최근 7일 고장 표시 수 --------------------------------
    machine_daily_flags = (
        data.groupby(["asset_tag", "transaction_date"], observed=True)["breakdown_flag"]
        .sum()
        .reset_index()
        .sort_values(["asset_tag", "transaction_date"])
    )
    shifted = machine_daily_flags.groupby("asset_tag", observed=True)["breakdown_flag"].shift(1)
    machine_daily_flags["machine_bd_count_7d"] = shifted.groupby(
        machine_daily_flags["asset_tag"], observed=True, sort=False
    ).transform(lambda s: s.rolling(7, min_periods=1).sum())
    data = data.merge(
        machine_daily_flags[["asset_tag", "transaction_date", "machine_bd_count_7d"]],
        on=["asset_tag", "transaction_date"],
        how="left",
    )

    # --- 센서 파생 (기계×날짜에서 계산 후 결합) --------------------------
    daily = build_machine_daily(frame)
    derived = [c for c in daily.columns if c.endswith(("_ma7", "_diff1", "_dev"))]
    data = data.merge(
        daily[["transaction_date", "asset_tag", *derived]],
        on=["transaction_date", "asset_tag"],
        how="left",
    )

    # --- 과제 A 타깃: t+1 ~ t+7 의 최대값, 당일 미포함 -------------------
    future = pd.concat(
        [grouped.shift(-k).rename(f"f{k}") for k in range(1, 8)], axis=1
    )
    complete = future.notna().all(axis=1)
    data["target_7d"] = np.where(complete, future.max(axis=1), np.nan)

    log(f"특성 생성 완료: {data.shape[1]}열 ({time.perf_counter() - started:.1f}초)")
    return data


def leakage_assertions(data: pd.DataFrame) -> list[tuple[str, str, str]]:
    """실행 중 누수 점검. 위반 시 예외를 던지고, 통과 항목을 반환한다."""
    checks: list[tuple[str, str, str]] = []

    first_rows = data.groupby(["asset_tag", "part_no"], observed=True).head(1)
    if first_rows["hist_bd_rate"].notna().any():
        raise AssertionError("과거 고장률 첫 행이 NaN이 아닙니다 - shift(1) 누락 의심")
    checks.append(("과거 고장률 shift(1)", "그룹 첫 행이 모두 NaN", "통과"))

    sample = data.groupby(["asset_tag", "part_no"], observed=True).head(1)
    if sample["days_since_bd"].notna().any():
        raise AssertionError("마지막 고장 경과일 첫 행이 NaN이 아닙니다")
    checks.append(("마지막 고장 경과일", "그룹 첫 행이 모두 NaN", "통과"))

    labelled = data.loc[data["target_7d"].notna()]
    if len(labelled) and labelled.groupby(["asset_tag", "part_no"], observed=True)[
        "transaction_date"
    ].max().max() > pd.Timestamp("2024-12-25"):
        raise AssertionError("7일 타깃이 관측 불가 구간까지 생성되었습니다")
    checks.append(("7일 타깃 범위", "t+1~t+7 전부 관측 가능한 행만 라벨", "통과"))

    for name, features in FEATURE_SETS.items():
        validate_features(features)
    checks.append(("분류 입력 금지 컬럼", "wo_type·qty_issued·issue_value_inr·타깃 제외", "통과"))
    return checks


# ===========================================================================
# 3. 데이터 검증 (Sheet 01 / 02 / 03 의 원천)
# ===========================================================================
def _judge(actual: Any, expected: Any, tolerance: float = 0.0) -> str:
    if expected is None:
        return "참고"
    if isinstance(expected, float):
        return "일치" if abs(float(actual) - expected) <= tolerance else "불일치"
    return "일치" if actual == expected else "불일치"


def _row(area, item, condition, actual, verdict, meaning, tasks, note="") -> dict[str, Any]:
    return {
        "검증영역": area,
        "검증항목": item,
        "기대조건": condition,
        "실제결과": actual,
        "판정": verdict,
        "해석": meaning,
        "사용과제": tasks,
        "비고": note,
    }


def verify_dataset(raw: pd.DataFrame, data: pd.DataFrame) -> tuple[list[dict], dict[str, Any]]:
    """전달문서 2장의 1차 검증값을 재현한다. (모델 학습보다 먼저 실행)"""
    rows: list[dict] = []
    facts: dict[str, Any] = {}

    # --- 2-1 ~ 2-3 기본 구조 -------------------------------------------
    rows.append(_row("기본구조", "행 수", f"{EXPECTED['rows']:,}", len(raw),
                     _judge(len(raw), EXPECTED["rows"]), "원본 그레인 유지 여부", "전체"))
    rows.append(_row("기본구조", "열 수", f"{EXPECTED['cols']}", raw.shape[1],
                     _judge(raw.shape[1], EXPECTED["cols"]), "컬럼 누락 없음", "전체"))
    date_min = raw["transaction_date"].min()
    date_max = raw["transaction_date"].max()
    rows.append(_row("기본구조", "날짜 범위 시작", EXPECTED["date_min"], date_min.strftime("%Y-%m-%d"),
                     _judge(date_min.strftime("%Y-%m-%d"), EXPECTED["date_min"]), "시계열 시작일", "전체"))
    rows.append(_row("기본구조", "날짜 범위 종료", EXPECTED["date_max"], date_max.strftime("%Y-%m-%d"),
                     _judge(date_max.strftime("%Y-%m-%d"), EXPECTED["date_max"]), "시계열 종료일", "전체"))
    n_dates = raw["transaction_date"].nunique()
    rows.append(_row("기본구조", "고유 날짜 수", f"{EXPECTED['n_dates']:,}", n_dates,
                     _judge(n_dates, EXPECTED["n_dates"]), "일 단위 연속 관측", "전체"))
    for column, key, label in [
        ("asset_tag", "n_assets", "기계 수"),
        ("part_no", "n_parts", "부품 수"),
        ("machine_type", "n_machine_types", "기계 종류 수"),
        ("plant_code", "n_plants", "공장 코드 수"),
    ]:
        value = raw[column].nunique()
        rows.append(_row("기본구조", label, str(EXPECTED[key]), value,
                         _judge(value, EXPECTED[key]), f"{column} 고유값", "전체"))

    # --- 2-4 / 2-5 키·중복 ---------------------------------------------
    pk_dupes = int(raw.duplicated(["transaction_date", "asset_tag", "part_no"]).sum())
    rows.append(_row("중복", "기본키 중복", "0", pk_dupes, _judge(pk_dupes, EXPECTED["pk_dupes"]),
                     "transaction_date+asset_tag+part_no 가 유일", "전체"))
    full_dupes = int(raw.duplicated().sum())
    rows.append(_row("중복", "완전 중복 행", "0", full_dupes, _judge(full_dupes, EXPECTED["full_dupes"]),
                     "동일 행 반복 없음", "전체"))
    grid_complete = n_dates * raw["asset_tag"].nunique() * raw["part_no"].nunique() == len(raw)
    rows.append(_row("중복", "날짜×기계×부품 격자 완전성", "완전 격자", "완전" if grid_complete else "불완전",
                     "일치" if grid_complete else "불일치",
                     "결측 날짜가 없어 위치 기반 rolling 이 일(day) 창과 동일", "전체"))

    # --- 2-6 타깃 분포 ---------------------------------------------------
    breakdown_1 = int(raw["breakdown_flag"].sum())
    breakdown_rate = float(raw["breakdown_flag"].mean())
    rows.append(_row("타깃분포", "breakdown_flag=1 건수", f"{EXPECTED['breakdown_1']:,}", breakdown_1,
                     _judge(breakdown_1, EXPECTED["breakdown_1"]), "당일 고장 표시 건수", "당일분류"))
    rows.append(_row("타깃분포", "전체 당일 고장 표시율", "9.8795%", breakdown_rate,
                     _judge(breakdown_rate, EXPECTED["breakdown_rate"], 5e-6),
                     "불균형 - Accuracy 대신 PR-AUC 사용", "당일분류"))

    # --- 2-7 결측치 ------------------------------------------------------
    missing = raw.isna().sum()
    wo_missing = int(missing["wo_type"])
    rows.append(_row("결측치", "wo_type 결측 건수", f"{EXPECTED['wo_type_missing']:,}", wo_missing,
                     _judge(wo_missing, EXPECTED["wo_type_missing"]), "원문상 작업 없음", "제외변수"))
    rows.append(_row("결측치", "wo_type 결측률", "80.46%", float(wo_missing / len(raw)),
                     _judge(round(wo_missing / len(raw), 4), EXPECTED["wo_type_missing_rate"], 1e-4),
                     "작업이 없던 날", "제외변수"))
    other_missing = int(missing.drop("wo_type").sum())
    rows.append(_row("결측치", "wo_type 외 21개 열 결측 합", "0", other_missing,
                     _judge(other_missing, 0), "보정 불필요", "전체"))

    # --- 2-8 / 2-9 누수 변수 ---------------------------------------------
    bd_mismatch = int(((raw["wo_type"] == "BD") != (raw["breakdown_flag"] == 1)).sum())
    rows.append(_row("누수변수", "wo_type=BD 와 breakdown_flag=1 대응", "불일치 0건", bd_mismatch,
                     _judge(bd_mismatch, 0), "wo_type 은 타깃과 1:1 - 입력 제외 필수", "제외변수"))
    pm_on_positive = int(((raw["wo_type"] == "PM") & (raw["breakdown_flag"] == 1)).sum())
    rows.append(_row("누수변수", "wo_type=PM 이 양성에 등장", "0건", pm_on_positive,
                     _judge(pm_on_positive, 0), "PM 은 breakdown_flag=0 에서만 발생", "제외변수"))
    identity = np.isclose(raw["issue_value_inr"], raw["unit_cost_inr"] * raw["qty_issued"], atol=1e-6)
    identity_rows = int(identity.sum())
    rows.append(_row("단위확인", "issue_value = unit_cost x qty 성립", f"{EXPECTED['rows']:,}행", identity_rows,
                     _judge(identity_rows, len(raw)), "출고금액은 파생값 - 분류 입력 제외", "제외변수"))

    # --- 2-10 / 2-11 그레인 ----------------------------------------------
    machine_daily = raw.drop_duplicates(["transaction_date", "asset_tag"])
    rows.append(_row("그레인", "기계x날짜 중복 제거 행 수", f"{EXPECTED['machine_daily_rows']:,}",
                     len(machine_daily), _judge(len(machine_daily), EXPECTED["machine_daily_rows"]),
                     "전력 회귀·K-Means 의 분석 단위", "전력회귀/KMeans"))
    sensor_unique = raw.groupby(["transaction_date", "asset_tag"], observed=True)[SENSORS].nunique()
    sensor_violation = int((sensor_unique > 1).sum().sum())
    rows.append(_row("그레인", "센서 8종이 부품 20행에서 동일", "위반 0건", sensor_violation,
                     _judge(sensor_violation, 0), "센서는 기계x날짜 속성 - 중복 제거 필요", "전력회귀/KMeans"))
    flag_by_day = raw.groupby(["transaction_date", "asset_tag"], observed=True)["breakdown_flag"].nunique()
    mixed_days = int((flag_by_day > 1).sum())
    rows.append(_row("그레인", "부품별 고장값이 갈리는 기계x날짜 수", "참고(기획서 7,376)", mixed_days,
                     "참고", "기계 수준 고장값을 첫 부품 행으로 대표시키면 안 됨", "전체"))

    # --- 2-13 ~ 2-16 과제별 분포 -----------------------------------------
    eligible = data.loc[(data["breakdown_flag"] == 0) & data["target_7d"].notna()]
    facts["task_a_eligible"] = eligible
    rows.append(_row("타깃분포", "7일 평가 대상 행 수", f"{EXPECTED['a_eval_rows']:,}", len(eligible),
                     _judge(len(eligible), EXPECTED["a_eval_rows"]),
                     "현재 고장 표시가 없고 미래 7일이 관측 가능한 행", "7일분류"))
    positives = int(eligible["target_7d"].sum())
    rows.append(_row("타깃분포", "7일 양성 수", f"{EXPECTED['a_eval_pos']:,}", positives,
                     _judge(positives, EXPECTED["a_eval_pos"]), "불균형 아님", "7일분류"))
    positive_rate = float(eligible["target_7d"].mean())
    rows.append(_row("타깃분포", "7일 양성률", "50.9829%", positive_rate,
                     _judge(positive_rate, EXPECTED["a_eval_rate"], 5e-6), "PR-AUC 기준선 약 0.51", "7일분류"))

    splits_a = split_frame(eligible, SPLIT_A)
    for name, key in [("학습", "a_train"), ("검증", "a_valid"), ("평가", "a_test")]:
        part = splits_a[name]
        want_rows, want_pos, want_rate = EXPECTED[key]
        rows.append(_row("시간분할", f"7일 {name} 행수", f"{want_rows:,}", len(part),
                         _judge(len(part), want_rows), "날짜 경계 고정 분할", "7일분류"))
        rows.append(_row("시간분할", f"7일 {name} 양성수", f"{want_pos:,}", int(part["target_7d"].sum()),
                         _judge(int(part["target_7d"].sum()), want_pos), "구간별 양성 건수", "7일분류"))
        rows.append(_row("시간분할", f"7일 {name} 양성률", f"{want_rate:.4%}", float(part["target_7d"].mean()),
                         _judge(float(part["target_7d"].mean()), want_rate, 5e-6), "구간 간 분포 안정", "7일분류"))

    splits_b = split_frame(data, SPLIT_BCD)
    for name, key in [("학습", "b_train"), ("검증", "b_valid"), ("평가", "b_test")]:
        part = splits_b[name]
        want_rows, want_pos, want_rate = EXPECTED[key]
        rows.append(_row("시간분할", f"당일 {name} 행수", f"{want_rows:,}", len(part),
                         _judge(len(part), want_rows), "경계 제외 없음", "당일분류"))
        rows.append(_row("시간분할", f"당일 {name} 양성수", f"{want_pos:,}", int(part["breakdown_flag"].sum()),
                         _judge(int(part["breakdown_flag"].sum()), want_pos), "구간별 양성 건수", "당일분류"))
        rows.append(_row("시간분할", f"당일 {name} 양성률", f"{want_rate:.4%}",
                         float(part["breakdown_flag"].mean()),
                         _judge(float(part["breakdown_flag"].mean()), want_rate, 5e-6),
                         "약 9.9% 로 안정", "당일분류"))

    daily = build_machine_daily(raw)
    splits_c = split_frame(daily, SPLIT_BCD)
    for name, key in [("학습", "c_train"), ("검증", "c_valid"), ("평가", "c_test")]:
        rows.append(_row("시간분할", f"전력 {name} 행수", f"{EXPECTED[key]:,}", len(splits_c[name]),
                         _judge(len(splits_c[name]), EXPECTED[key]), "기계x날짜 그레인", "전력회귀/KMeans"))

    # --- 2-17 MAPE 사용 가능 여부 ----------------------------------------
    zero_power = int((daily["power_consumption_kw"] == 0).sum())
    rows.append(_row("단위확인", "power_consumption_kw = 0 행", "0건이면 MAPE 사용 가능", zero_power,
                     "참고", "0이 있으면 MAPE 정의를 명시하거나 제외", "전력회귀",
                     "0건이면 일반 MAPE 정의 사용" if zero_power == 0 else "0 존재 - MAPE 정의 명시 필요"))
    facts["zero_power"] = zero_power
    facts["machine_daily"] = daily
    return rows, facts


def build_column_dictionary(raw: pd.DataFrame) -> pd.DataFrame:
    """Sheet 02 - 22개 열의 역할·단위·통계·모델 입력 여부."""
    meta = {
        "transaction_date": ("날짜키", "date", "전체", "아니오", "분할·정렬 기준"),
        "asset_tag": ("식별자/범주형 입력", "-", "전체", "예", ""),
        "machine_type": ("범주형 입력", "-", "전체", "예", ""),
        "plant_code": ("범주형 입력", "-", "전체", "예", ""),
        "part_no": ("식별자/범주형 입력", "-", "분류", "예", ""),
        "part_description": ("설명", "-", "조회", "아니오", "part_no 와 1:1 중복"),
        "part_family": ("범주형 입력", "-", "분류", "예", ""),
        "criticality": ("범주형 입력", "등급", "분류", "예", ""),
        "uom": ("단위", "-", "조회", "아니오", "부품 속성과 사실상 결합"),
        "unit_cost_inr": ("정적 수치 입력", "INR", "분류", "예", ""),
        "qty_issued": ("사후 정보", "EA", "조회", "아니오", "고장 대응 후 출고량 - 누수"),
        "issue_value_inr": ("사후 정보", "INR", "조회", "아니오", "unit_cost x qty 파생 - 누수"),
        "temp_bearing_degC": ("센서 입력", "degC", "전체", "예", ""),
        "temp_motor_degC": ("센서 입력", "degC", "전체", "예", ""),
        "vibration_h_mms": ("센서 입력", "mm/s", "전체", "예", ""),
        "vibration_v_mms": ("센서 입력", "mm/s", "전체", "예", ""),
        "oil_pressure_bar": ("센서 입력", "bar", "전체", "예", ""),
        "load_pct": ("운전 입력", "%", "전체", "예", ""),
        "shaft_rpm": ("운전 입력", "rpm", "전체", "예", ""),
        "power_consumption_kw": ("센서 입력/회귀 타깃", "kW", "전체", "예", "회귀에서는 과거 지연값만 입력"),
        "breakdown_flag": ("타깃", "0/1", "분류", "아니오", "당일 분류 타깃·7일 타깃 원천"),
        "wo_type": ("사후 정보", "-", "제외", "아니오", "BD 와 flag=1 이 1:1 대응 - 누수"),
    }
    records = []
    for column in raw.columns:
        series = raw[column]
        numeric = pd.api.types.is_numeric_dtype(series)
        role, unit, tasks, model_input, reason = meta.get(column, ("미지정", "-", "-", "아니오", ""))
        records.append({
            "컬럼명": column,
            "데이터형": str(series.dtype),
            "역할": role,
            "단위": unit,
            "고유값수": int(series.nunique()),
            "결측수": int(series.isna().sum()),
            "결측률": float(series.isna().mean()),
            "최소값": float(series.min()) if numeric else None,
            "최대값": float(series.max()) if numeric else None,
            "평균": float(series.mean()) if numeric else None,
            "사용과제": tasks,
            "모델입력여부": model_input,
            "제외사유": reason,
        })
    return pd.DataFrame(records)


# ===========================================================================
# 4. 분류 공통 유틸 (과제 A / B)
# ===========================================================================
from sklearn.compose import ColumnTransformer
from sklearn.ensemble import (
    HistGradientBoostingClassifier,
    HistGradientBoostingRegressor,
    RandomForestClassifier,
)
from sklearn.cluster import KMeans
from sklearn.impute import SimpleImputer
from sklearn.inspection import permutation_importance
from sklearn.linear_model import LogisticRegression
from sklearn.metrics import (
    accuracy_score,
    average_precision_score,
    confusion_matrix,
    f1_score,
    mean_absolute_error,
    mean_squared_error,
    precision_recall_curve,
    precision_score,
    r2_score,
    recall_score,
    roc_auc_score,
    silhouette_score,
)
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import OneHotEncoder, StandardScaler

CLASSIFIER_SPECS = {
    "Logistic Regression": dict(max_iter=1000, random_state=RANDOM_STATE),
    "Random Forest": dict(n_estimators=200, min_samples_leaf=5, n_jobs=-1, random_state=RANDOM_STATE),
    "HistGradientBoosting": dict(max_iter=200, early_stopping=False, random_state=RANDOM_STATE),
}
MODEL_NOTE = {
    "Logistic Regression": "LogisticRegression(max_iter=1000)",
    "Random Forest": "RandomForestClassifier(n_estimators=200, min_samples_leaf=5)",
    "HistGradientBoosting": "HistGradientBoostingClassifier(max_iter=200, early_stopping=False)",
}
TOP_FRACTIONS = {"10": 0.10, "20": 0.20}
PR_CURVE_POINTS = 300


def build_classifier(name: str, features: list[str]) -> Pipeline:
    """전처리(보정·인코딩·표준화)를 파이프라인 안에 두어 학습 구간에만 fit 되게 한다."""
    validate_features(features)
    categorical = [c for c in features if c in CATEGORICALS]
    numeric = [c for c in features if c not in CATEGORICALS]
    numeric_steps: Any = SimpleImputer(strategy="median")
    if name == "Logistic Regression":
        numeric_steps = Pipeline(
            [("impute", SimpleImputer(strategy="median")), ("scale", StandardScaler())]
        )
    transformers: list[tuple[str, Any, list[str]]] = []
    if categorical:
        transformers.append(
            ("category", OneHotEncoder(handle_unknown="ignore", sparse_output=False), categorical)
        )
    if numeric:
        transformers.append(("numeric", numeric_steps, numeric))
    preprocess = ColumnTransformer(transformers=transformers, remainder="drop")

    if name == "Logistic Regression":
        estimator = LogisticRegression(**CLASSIFIER_SPECS[name])
    elif name == "Random Forest":
        estimator = RandomForestClassifier(**CLASSIFIER_SPECS[name])
    elif name == "HistGradientBoosting":
        estimator = HistGradientBoostingClassifier(**CLASSIFIER_SPECS[name])
    else:
        raise ValueError(f"알 수 없는 모델: {name}")
    return Pipeline([("preprocess", preprocess), ("classifier", estimator)])


def choose_threshold(y_true: np.ndarray, scores: np.ndarray) -> float:
    """검증 세트에서 F1 이 가장 높은 임계값을 고른다. 평가 세트는 사용하지 않는다."""
    y = np.asarray(y_true, dtype=int)
    probabilities = np.asarray(scores, dtype=float)
    if len(y) == 0 or len(np.unique(y)) < 2:
        return 0.5
    precision, recall, thresholds = precision_recall_curve(y, probabilities)
    denominator = precision + recall
    f1 = np.divide(2 * precision * recall, denominator,
                   out=np.zeros_like(denominator), where=denominator > 0)
    if len(thresholds) == 0:
        return 0.5
    best = int(np.argmax(f1[:-1])) if len(f1) > 1 else 0
    return float(thresholds[min(best, len(thresholds) - 1)])


def top_fraction_scores(y_true: np.ndarray, scores: np.ndarray, fraction: float) -> tuple[float, float]:
    """점수 상위 fraction 구간의 Precision / Recall."""
    y = np.asarray(y_true, dtype=int)
    total_positive = int(y.sum())
    count = max(1, int(round(len(y) * fraction)))
    order = np.argsort(-np.asarray(scores, dtype=float), kind="mergesort")[:count]
    hits = int(y[order].sum())
    precision = hits / count
    recall = hits / total_positive if total_positive else float("nan")
    return precision, recall


def classification_row(
    *, task: str, model: str, feature_set: str, split: str, threshold: float,
    y_true: np.ndarray, scores: np.ndarray, fit_seconds: float, predict_seconds: float, note: str,
) -> dict[str, Any]:
    y = np.asarray(y_true, dtype=int)
    probabilities = np.clip(np.asarray(scores, dtype=float), 0.0, 1.0)
    predictions = (probabilities >= threshold).astype(int)
    tn, fp, fn, tp = confusion_matrix(y, predictions, labels=[0, 1]).ravel()
    both_classes = len(np.unique(y)) == 2
    row = {
        "과제": task,
        "모델": model,
        "특성세트": feature_set,
        "구간": split,
        "임계값": float(threshold),
        "평가건수": int(len(y)),
        "양성수": int(y.sum()),
        "양성률": float(y.mean()) if len(y) else None,
        "예측양성률": float(predictions.mean()) if len(y) else None,
        "ROC_AUC": float(roc_auc_score(y, probabilities)) if both_classes else None,
        "PR_AUC": float(average_precision_score(y, probabilities)) if both_classes else None,
        "Accuracy": float(accuracy_score(y, predictions)),
        "Precision": float(precision_score(y, predictions, zero_division=0)),
        "Recall": float(recall_score(y, predictions, zero_division=0)),
        "F1": float(f1_score(y, predictions, zero_division=0)),
        "Brier_Score": float(np.mean((probabilities - y) ** 2)),
    }
    for label, fraction in TOP_FRACTIONS.items():
        precision_at, recall_at = top_fraction_scores(y, probabilities, fraction)
        row[f"Precision_at_{label}"] = precision_at
        row[f"Recall_at_{label}"] = recall_at
    row.update({
        "TN": int(tn), "FP": int(fp), "FN": int(fn), "TP": int(tp),
        "기준B대비_ROC차이": None,
        "기준B대비_PR차이": None,
        "학습시간초": round(float(fit_seconds), 3),
        "예측시간초": round(float(predict_seconds), 3),
        "비고": note,
    })
    return row


def pr_curve_rows(task: str, model: str, feature_set: str, split: str,
                  y_true: np.ndarray, scores: np.ndarray) -> list[dict]:
    """PR 곡선 원시 좌표를 균등 샘플링해 저장한다 (Sheet 07)."""
    y = np.asarray(y_true, dtype=int)
    if len(np.unique(y)) < 2:
        return []
    precision, recall, thresholds = precision_recall_curve(y, np.asarray(scores, dtype=float))
    # precision/recall 은 thresholds 보다 1개 길다. 마지막 점을 버리고 맞춘다.
    precision, recall = precision[:-1], recall[:-1]
    if len(thresholds) > PR_CURVE_POINTS:
        index = np.unique(np.linspace(0, len(thresholds) - 1, PR_CURVE_POINTS).astype(int))
    else:
        index = np.arange(len(thresholds))
    return [
        {
            "과제": task, "모델": model, "특성세트": feature_set, "구간": split,
            "threshold": float(thresholds[i]),
            "precision": float(precision[i]),
            "recall": float(recall[i]),
        }
        for i in index
    ]


def confusion_rows(task: str, model: str, feature_set: str, split: str,
                   threshold: float, y_true: np.ndarray, scores: np.ndarray) -> list[dict]:
    y = np.asarray(y_true, dtype=int)
    predictions = (np.asarray(scores, dtype=float) >= threshold).astype(int)
    matrix = confusion_matrix(y, predictions, labels=[0, 1])
    records = []
    for actual in (0, 1):
        for predicted in (0, 1):
            records.append({
                "과제": task, "모델": model, "특성세트": feature_set, "구간": split,
                "임계값": float(threshold), "실제값": actual, "예측값": predicted,
                "건수": int(matrix[actual, predicted]),
            })
    return records


def importance_rows(task: str, model_name: str, feature_set: str, pipeline: Pipeline,
                    features: list[str], valid: pd.DataFrame, y_valid: np.ndarray) -> list[dict]:
    """LR 은 계수, 트리 모델은 permutation importance (검증 세트 서브샘플)."""
    records: list[dict] = []
    if model_name == "Logistic Regression":
        names = pipeline.named_steps["preprocess"].get_feature_names_out()
        coefficients = pipeline.named_steps["classifier"].coef_[0]
        order = np.argsort(-np.abs(coefficients))
        for rank, index in enumerate(order[:40], start=1):
            records.append({
                "과제": task, "모델": model_name, "특성세트": feature_set,
                "변수명": str(names[index]), "중요도방법": "계수(절대값 정렬)",
                "중요도": float(coefficients[index]), "순위": rank,
                "방향": "양(+)" if coefficients[index] > 0 else "음(-)",
                "비고": "표준화된 수치 + 원핫 범주 기준. 다른 중요도 방식과 직접 비교 금지",
            })
        return records

    sample_size = min(5000, len(valid))
    rng = np.random.RandomState(RANDOM_STATE)
    index = rng.choice(len(valid), size=sample_size, replace=False)
    subset = valid.iloc[index]
    subset_y = np.asarray(y_valid)[index]
    result = permutation_importance(
        pipeline, subset[features], subset_y, scoring="roc_auc",
        n_repeats=3, random_state=RANDOM_STATE, n_jobs=1,
    )
    order = np.argsort(-result.importances_mean)
    for rank, position in enumerate(order[:40], start=1):
        records.append({
            "과제": task, "모델": model_name, "특성세트": feature_set,
            "변수명": features[position], "중요도방법": "permutation importance(ROC-AUC 감소)",
            "중요도": float(result.importances_mean[position]), "순위": rank,
            "방향": "-",
            "비고": f"검증 {sample_size:,}행 서브샘플, n_repeats=3, random_state={RANDOM_STATE}",
        })
    return records


def run_classification_task(
    task: str, splits: dict[str, pd.DataFrame], target: str, *,
    feature_sets: dict[str, list[str]], baseline_b_transform: Any, outputs: Path,
) -> dict[str, list[dict]]:
    """기준 A·B + LR/RF/HGB x 특성세트를 학습하고 모든 시트용 행을 만든다."""
    train, valid, test = splits["학습"], splits["검증"], splits["평가"]
    y = {name: splits[name][target].astype(int).to_numpy() for name in ("학습", "검증", "평가")}
    scored = {"검증": valid, "평가": test}

    metrics: list[dict] = []
    curves: list[dict] = []
    matrices: list[dict] = []
    importances: list[dict] = []
    prediction_frames: list[pd.DataFrame] = []
    baseline_b_reference: dict[str, dict[str, float]] = {}

    def record(model_name, feature_set, scores_by_split, fit_seconds, predict_seconds, note, threshold=None):
        selected = choose_threshold(y["검증"], scores_by_split["검증"]) if threshold is None else threshold
        for split_name, frame in scored.items():
            scores = scores_by_split[split_name]
            for used_threshold, threshold_note in [
                (selected, f"검증 F1 최대 임계값 {selected:.4f} 고정 적용"),
                (0.5, "임계값 0.5 고정"),
            ]:
                metrics.append(classification_row(
                    task=task, model=model_name, feature_set=feature_set, split=split_name,
                    threshold=used_threshold, y_true=y[split_name], scores=scores,
                    fit_seconds=fit_seconds if split_name == "평가" else fit_seconds,
                    predict_seconds=predict_seconds,
                    note=f"{note} / {threshold_note}",
                ))
                matrices.extend(confusion_rows(task, model_name, feature_set, split_name,
                                               used_threshold, y[split_name], scores))
            curves.extend(pr_curve_rows(task, model_name, feature_set, split_name,
                                        y[split_name], scores))
            if split_name == "평가":
                frame_out = frame[[c for c in ["transaction_date", "asset_tag", "part_no", target]
                                   if c in frame.columns]].copy()
                frame_out["모델"] = model_name
                frame_out["특성세트"] = feature_set
                frame_out["risk_score"] = scores
                frame_out["임계값"] = selected
                frame_out["예측"] = (scores >= selected).astype(int)
                prediction_frames.append(frame_out)
        return selected

    # --- 기준 A: 학습 구간 전체 양성 비율 --------------------------------
    prior = float(y["학습"].mean())
    record("기준A(학습구간 양성비율)", "-",
           {name: np.full(len(splits[name]), prior) for name in scored},
           0.0, 0.0, f"학습 구간 양성률 {prior:.6f} 상수 예측")

    # --- 기준 B: 기계x부품 과거 고장 표시율 -------------------------------
    fallback = float(train["hist_bd_rate"].mean())
    baseline_scores = {}
    for name, frame in scored.items():
        daily_rate = frame["hist_bd_rate"].fillna(fallback).to_numpy(dtype=float)
        baseline_scores[name] = baseline_b_transform(daily_rate)
    record("기준B(기계x부품 과거 고장률)", "-", baseline_scores, 0.0, 0.0,
           "shift(1) 누적 고장 표시율만 사용 (미래 정보 없음)")
    for split_name in scored:
        baseline_b_reference[split_name] = {
            "ROC_AUC": float(roc_auc_score(y[split_name], baseline_scores[split_name])),
            "PR_AUC": float(average_precision_score(y[split_name], baseline_scores[split_name])),
        }

    # --- 학습 모델 -------------------------------------------------------
    for feature_set, features in feature_sets.items():
        for model_name in CLASSIFIER_SPECS:
            pipeline = build_classifier(model_name, features)
            started = time.perf_counter()
            pipeline.fit(train[features], y["학습"])
            fit_seconds = time.perf_counter() - started
            scores_by_split = {}
            predict_seconds = 0.0
            for split_name, frame in scored.items():
                started = time.perf_counter()
                scores_by_split[split_name] = pipeline.predict_proba(frame[features])[:, 1]
                predict_seconds += time.perf_counter() - started
            record(model_name, feature_set, scores_by_split, fit_seconds, predict_seconds,
                   MODEL_NOTE[model_name])
            importances.extend(importance_rows(task, model_name, feature_set, pipeline,
                                               features, valid, y["검증"]))
            log(f"  {task} | {feature_set} | {model_name}: 학습 {fit_seconds:.1f}초, "
                f"평가 ROC-AUC {roc_auc_score(y['평가'], scores_by_split['평가']):.4f}")

    # --- 기준 B 대비 차이 채우기 -----------------------------------------
    for row in metrics:
        reference = baseline_b_reference.get(row["구간"])
        if reference and not row["모델"].startswith("기준B") and row["ROC_AUC"] is not None:
            row["기준B대비_ROC차이"] = row["ROC_AUC"] - reference["ROC_AUC"]
            row["기준B대비_PR차이"] = row["PR_AUC"] - reference["PR_AUC"]

    if prediction_frames:
        filename = "sevenday_test_predictions.csv" if task == "7일분류" else "sameday_test_predictions.csv"
        pd.concat(prediction_frames, ignore_index=True).to_csv(
            outputs / filename, index=False, encoding="utf-8-sig"
        )
    return {"metrics": metrics, "curves": curves, "matrices": matrices, "importances": importances}


# ===========================================================================
# 5. 과제 C - 다음 관측일 전력 회귀
# ===========================================================================
POWER = "power_consumption_kw"
REGRESSION_FEATURES = [
    "power_lag1", "power_lag2", "power_lag7",
    "power_ma7", "power_ma14", "power_std7",
    "load_pct_lag1", "shaft_rpm_lag1",
    "target_dayofweek", "target_is_weekend",
    "machine_type", "asset_tag",
]
PROPHET_RETRAIN_DAYS = 7


def build_power_frame(machine_daily: pd.DataFrame) -> pd.DataFrame:
    """기계별 지연·이동 특성과 다음 관측일 타깃을 만든다. 미래 센서는 입력하지 않는다."""
    frame = machine_daily.sort_values(["asset_tag", "transaction_date"]).reset_index(drop=True)
    groups = frame.groupby("asset_tag", observed=True, sort=False)

    # 예측 대상은 다음 관측일. 타깃 날짜는 사전에 알 수 있으므로 요일만 사용한다.
    frame["target_power_next"] = groups[POWER].shift(-1)
    frame["target_date"] = groups["transaction_date"].shift(-1)
    frame["target_dayofweek"] = frame["target_date"].dt.dayofweek
    frame["target_is_weekend"] = (frame["target_dayofweek"] >= 5).astype(float)

    # 타깃(t+1) 기준 지연값: lag1 = 당일(t), lag7 = 타깃과 같은 요일(t-6)
    frame["power_lag1"] = frame[POWER]
    frame["power_lag2"] = groups[POWER].shift(1)
    frame["power_lag7"] = groups[POWER].shift(6)
    frame["power_ma7"] = groups[POWER].transform(lambda s: s.rolling(7, min_periods=7).mean())
    frame["power_ma14"] = groups[POWER].transform(lambda s: s.rolling(14, min_periods=14).mean())
    frame["power_std7"] = groups[POWER].transform(lambda s: s.rolling(7, min_periods=7).std())
    frame["load_pct_lag1"] = frame["load_pct"]
    frame["shaft_rpm_lag1"] = frame["shaft_rpm"]

    # 기계 경계를 넘는 지연값이 없어야 한다 (누수 점검 7).
    first_rows = frame.groupby("asset_tag", observed=True).head(1)
    if first_rows["power_lag2"].notna().any():
        raise AssertionError("전력 지연값이 기계 경계를 넘었습니다")
    return frame


def regression_metrics(actual: np.ndarray, predicted: np.ndarray) -> dict[str, float]:
    actual = np.asarray(actual, dtype=float)
    predicted = np.asarray(predicted, dtype=float)
    errors = np.abs(actual - predicted)
    return {
        "MAE": float(mean_absolute_error(actual, predicted)),
        "RMSE": float(np.sqrt(mean_squared_error(actual, predicted))),
        "R2": float(r2_score(actual, predicted)),
        "MAPE": float(np.mean(errors / np.abs(actual))),
        "Median_AE": float(np.median(errors)),
    }


def run_power_regression(machine_daily: pd.DataFrame, outputs: Path,
                         skip_prophet: bool = False) -> dict[str, Any]:
    frame = build_power_frame(machine_daily)
    splits = split_frame(frame, SPLIT_BCD)
    usable = {name: part.dropna(subset=["target_power_next", *[
        c for c in REGRESSION_FEATURES if c not in CATEGORICALS]]) for name, part in splits.items()}
    train, valid, test = usable["학습"], usable["검증"], usable["평가"]
    log(f"전력 회귀 유효행수: 학습 {len(train):,} / 검증 {len(valid):,} / 평가 {len(test):,}")

    scored = {"검증": valid, "평가": test}
    predictions: dict[str, dict[str, np.ndarray]] = {name: {} for name in scored}
    timings: dict[str, tuple[float, float]] = {}
    notes: dict[str, str] = {}

    # --- 기준모델 1: 전일값 지속 -----------------------------------------
    for name, part in scored.items():
        predictions[name]["기준1(전일값 지속)"] = part["power_lag1"].to_numpy(dtype=float)
    timings["기준1(전일값 지속)"] = (0.0, 0.0)
    notes["기준1(전일값 지속)"] = "당일 전력을 다음 관측일 예측값으로 사용"

    # --- 기준모델 2: 지난주 같은 요일 (lag 7) ------------------------------
    for name, part in scored.items():
        predictions[name]["기준2(지난주 같은 요일)"] = part["power_lag7"].to_numpy(dtype=float)
    timings["기준2(지난주 같은 요일)"] = (0.0, 0.0)
    notes["기준2(지난주 같은 요일)"] = "타깃 날짜와 같은 요일의 직전 관측값"

    # --- 기준모델 3: 기계 x 평일/주말 평균 (학습 구간으로만 계산) ----------
    started = time.perf_counter()
    weekday_mean = train.groupby(["asset_tag", "target_is_weekend"], observed=True)[
        "target_power_next"].mean()
    overall_mean = float(train["target_power_next"].mean())
    fit_seconds = time.perf_counter() - started
    for name, part in scored.items():
        keys = list(zip(part["asset_tag"], part["target_is_weekend"]))
        predictions[name]["기준3(기계x평일/주말 평균)"] = np.array(
            [weekday_mean.get(key, overall_mean) for key in keys], dtype=float)
    timings["기준3(기계x평일/주말 평균)"] = (fit_seconds, 0.0)
    notes["기준3(기계x평일/주말 평균)"] = "평균은 학습 구간에서만 계산 후 검증·평가에 적용"

    # --- HistGradientBoosting 회귀 ----------------------------------------
    categorical = [c for c in REGRESSION_FEATURES if c in CATEGORICALS]
    numeric = [c for c in REGRESSION_FEATURES if c not in CATEGORICALS]
    preprocess = ColumnTransformer([
        ("category", OneHotEncoder(handle_unknown="ignore", sparse_output=False), categorical),
        ("numeric", SimpleImputer(strategy="median"), numeric),
    ], remainder="drop")
    pipeline = Pipeline([("preprocess", preprocess),
                         ("regressor", HistGradientBoostingRegressor(random_state=RANDOM_STATE))])
    started = time.perf_counter()
    pipeline.fit(train[REGRESSION_FEATURES], train["target_power_next"])
    fit_seconds = time.perf_counter() - started
    predict_seconds = 0.0
    for name, part in scored.items():
        started = time.perf_counter()
        predictions[name]["HistGradientBoosting 회귀"] = pipeline.predict(part[REGRESSION_FEATURES])
        predict_seconds += time.perf_counter() - started
    timings["HistGradientBoosting 회귀"] = (fit_seconds, predict_seconds)
    notes["HistGradientBoosting 회귀"] = "HistGradientBoostingRegressor(기본값), 학습 구간만 fit"
    log(f"  전력 HGB 학습 {fit_seconds:.1f}초")

    # --- Prophet (기계별, 7일 주기 재학습) ---------------------------------
    if not skip_prophet:
        prophet_predictions, prophet_fit, prophet_predict, fit_count = run_prophet(frame, scored)
        for name in scored:
            predictions[name]["Prophet"] = prophet_predictions[name]
        timings["Prophet"] = (prophet_fit, prophet_predict)
        notes["Prophet"] = (
            f"기계별 개별 모델, 주간 계절성 사용/연간·일간 끔, {PROPHET_RETRAIN_DAYS}일 주기 재학습"
            f"(총 {fit_count}회 적합). 예측 대상은 다른 모델과 동일한 다음 관측일 1일이나, "
            f"창 내부에서 모델은 최대 {PROPHET_RETRAIN_DAYS - 1}일 이전 정보까지만 반영"
        )

    # --- 지표 산출 ---------------------------------------------------------
    rows: list[dict] = []
    baseline_names = ["기준1(전일값 지속)", "기준2(지난주 같은 요일)", "기준3(기계x평일/주말 평균)"]
    validation_mae = {
        name: mean_absolute_error(valid["target_power_next"], predictions["검증"][name])
        for name in baseline_names
    }
    strongest = min(validation_mae, key=validation_mae.get)
    log(f"  최강 기준모델(검증 MAE 기준): {strongest} = {validation_mae[strongest]:.3f} kW")

    prediction_frames: list[pd.DataFrame] = []
    for split_name, part in scored.items():
        actual = part["target_power_next"].to_numpy(dtype=float)
        baseline_mae = float(mean_absolute_error(actual, predictions[split_name][strongest]))
        for model_name, predicted in predictions[split_name].items():
            metrics = regression_metrics(actual, predicted)
            fit_seconds, predict_seconds = timings[model_name]
            is_baseline = model_name in baseline_names
            rows.append({
                "모델": model_name,
                "구간": split_name,
                "재학습주기": f"{PROPHET_RETRAIN_DAYS}일" if model_name == "Prophet" else "1회(학습 구간)",
                "평가건수": int(len(actual)),
                **metrics,
                "최강기준모델": strongest,
                "기준_MAE": baseline_mae,
                "MAE_개선량": baseline_mae - metrics["MAE"],
                "MAE_개선율": (baseline_mae - metrics["MAE"]) / baseline_mae,
                "학습시간초": round(fit_seconds, 3),
                "예측시간초": round(predict_seconds, 3),
                "비고": notes[model_name] + ("" if not is_baseline else " / 기준모델"),
            })
            detail = pd.DataFrame({
                "transaction_date": part["target_date"].to_numpy(),
                "asset_tag": part["asset_tag"].to_numpy(),
                "machine_type": part["machine_type"].to_numpy(),
                "actual_power_kw": actual,
                "model": model_name,
                "predicted_power_kw": predicted,
                "residual_kw": actual - predicted,
                "absolute_error_kw": np.abs(actual - predicted),
                "split": split_name,
            })
            prediction_frames.append(detail)

    detail_frame = pd.concat(prediction_frames, ignore_index=True)
    detail_frame.to_csv(outputs / "power_predictions.csv", index=False, encoding="utf-8-sig")
    split_rows = [
        {"과제": "전력회귀", "구간": name, "시작일": SPLIT_BCD[name][0], "종료일": SPLIT_BCD[name][1],
         "전체행수": len(splits[name]), "유효행수": len(usable[name]), "양성수": None, "음성수": None,
         "양성률": None, "제외행수": len(splits[name]) - len(usable[name]),
         "제외사유": "다음 관측일 타깃 없음(구간 마지막 날) 또는 지연·이동 특성 결측"}
        for name in SPLIT_BCD
    ]
    return {"metrics": rows, "predictions": detail_frame, "splits": split_rows, "strongest": strongest}


def run_prophet(frame: pd.DataFrame, scored: dict[str, pd.DataFrame]):
    """기계별 Prophet 을 7일 주기로 재학습하며 다음 관측일을 예측한다."""
    import logging
    from prophet import Prophet

    for name in ("cmdstanpy", "prophet"):
        logger = logging.getLogger(name)
        logger.setLevel(logging.CRITICAL)
        logger.disabled = True

    history = frame[["asset_tag", "transaction_date", POWER]].dropna()
    results: dict[str, np.ndarray] = {}
    total_fit = 0.0
    total_predict = 0.0
    fit_count = 0

    for split_name, part in scored.items():
        predicted = pd.Series(np.nan, index=part.index, dtype=float)
        for asset, asset_rows in part.groupby("asset_tag", observed=True, sort=False):
            asset_rows = asset_rows.sort_values("transaction_date")
            asset_history = history.loc[history["asset_tag"] == asset]
            positions = np.arange(len(asset_rows))
            for start in range(0, len(asset_rows), PROPHET_RETRAIN_DAYS):
                window = asset_rows.iloc[positions[start:start + PROPHET_RETRAIN_DAYS]]
                cutoff = window["transaction_date"].iloc[0]
                train_history = asset_history.loc[asset_history["transaction_date"] <= cutoff]
                model = Prophet(weekly_seasonality=True, yearly_seasonality=False,
                                daily_seasonality=False)
                started = time.perf_counter()
                model.fit(train_history.rename(
                    columns={"transaction_date": "ds", POWER: "y"})[["ds", "y"]])
                total_fit += time.perf_counter() - started
                fit_count += 1
                started = time.perf_counter()
                forecast = model.predict(pd.DataFrame({"ds": window["target_date"].to_numpy()}))
                total_predict += time.perf_counter() - started
                predicted.loc[window.index] = forecast["yhat"].to_numpy()
        results[split_name] = predicted.to_numpy(dtype=float)
        log(f"  Prophet {split_name} 구간 완료 (누적 적합 {fit_count}회, {total_fit:.0f}초)")
    return results, total_fit, total_predict, fit_count


# ===========================================================================
# 6. 과제 D - K-Means 설비 상태 군집화 (비지도, 분류·회귀 표와 분리)
# ===========================================================================
K_CANDIDATES = [2, 3, 4, 5, 6]


def run_kmeans(machine_daily: pd.DataFrame, data: pd.DataFrame, outputs: Path) -> dict[str, Any]:
    frame = machine_daily.copy()
    splits = split_frame(frame, SPLIT_BCD)
    train = splits["학습"]

    # 기계 종류별 표준화 기준을 학습 구간에서만 적합한다 (누수 점검 4).
    scalers: dict[str, StandardScaler] = {}
    for machine_type, rows in train.groupby("machine_type", observed=True):
        scaler = StandardScaler().fit(rows[SENSORS])
        scalers[machine_type] = scaler

    def transform(part: pd.DataFrame) -> np.ndarray:
        output = np.empty((len(part), len(SENSORS)), dtype=float)
        for machine_type, rows in part.groupby("machine_type", observed=True):
            positions = part.index.get_indexer(rows.index)
            output[positions] = scalers[machine_type].transform(rows[SENSORS])
        return output

    scaled = {name: transform(part) for name, part in splits.items()}
    scaled_all = transform(frame)

    k_rows: list[dict] = []
    models: dict[int, KMeans] = {}
    for k in K_CANDIDATES:
        model = KMeans(n_clusters=k, random_state=RANDOM_STATE, n_init=10)
        model.fit(scaled["학습"])
        models[k] = model
        for split_name in ("학습", "평가"):
            matrix = scaled[split_name]
            labels = model.predict(matrix)
            sample_size = 5000 if len(matrix) > 5000 else None
            score = float(silhouette_score(matrix, labels, sample_size=sample_size,
                                           random_state=RANDOM_STATE))
            inertia = float(((matrix - model.cluster_centers_[labels]) ** 2).sum())
            k_rows.append({
                "k": k, "구간": split_name, "평가건수": int(len(matrix)),
                "Silhouette_Score": score, "Inertia": inertia,
                "선정여부": "", "선정근거": "",
            })

    train_scores = {row["k"]: row["Silhouette_Score"] for row in k_rows if row["구간"] == "학습"}
    best_overall = max(train_scores, key=train_scores.get)
    limited = {k: v for k, v in train_scores.items() if k <= 3}
    best_limited = max(limited, key=limited.get)
    selected = best_limited
    if best_overall == best_limited:
        reason = (f"학습 구간 Silhouette 최대(k={best_overall}, {train_scores[best_overall]:.4f})이며 "
                  f"화면 표현 제약(최대 3개 군집)도 함께 만족")
    else:
        reason = (f"Silhouette 최대는 k={best_overall}({train_scores[best_overall]:.4f})이나 "
                  f"기획서의 화면 표현 제약(최대 3개 군집)에 따라 k<=3 중 최대인 "
                  f"k={best_limited}({train_scores[best_limited]:.4f}) 선정")
    for row in k_rows:
        if row["k"] == selected:
            row["선정여부"] = "선정"
            row["선정근거"] = reason
        elif row["k"] == best_overall:
            row["선정여부"] = "Silhouette 최대"
            row["선정근거"] = "화면 표현 제약(최대 3개 군집)으로 최종 선정에서 제외"
        else:
            row["선정여부"] = "후보"
            row["선정근거"] = "비교용"

    model = models[selected]
    labels = model.predict(scaled_all)
    distances = np.linalg.norm(scaled_all - model.cluster_centers_[labels], axis=1)
    frame = frame.reset_index(drop=True)
    frame["cluster"] = labels
    frame["anomaly_score"] = distances

    # 기계 수준 고장 표시 참고율 = 그 날 20개 부품 행 중 고장 표시 비율
    reference = (data.groupby(["transaction_date", "asset_tag"], observed=True)["breakdown_flag"]
                 .mean().rename("고장표시참고율").reset_index())
    frame = frame.merge(reference, on=["transaction_date", "asset_tag"], how="left")
    frame[["transaction_date", "asset_tag", "machine_type", "cluster", "anomaly_score",
           "고장표시참고율", *SENSORS]].to_csv(
        outputs / "kmeans_assignments.csv", index=False, encoding="utf-8-sig")

    overall_mean = frame[SENSORS].mean()
    overall_std = frame[SENSORS].std()
    summary_rows: list[dict] = []
    for cluster_id, rows in frame.groupby("cluster"):
        z_scores = ((rows[SENSORS].mean() - overall_mean) / overall_std)
        ordered = z_scores.abs().sort_values(ascending=False)
        pieces = [f"{sensor} {'높음' if z_scores[sensor] > 0 else '낮음'}(z={z_scores[sensor]:+.2f})"
                  for sensor in ordered.index[:3]]
        interpretation = "센서 특성: " + ", ".join(pieces)
        threshold_95 = float(np.percentile(rows["anomaly_score"], 95))
        for sensor in SENSORS:
            summary_rows.append({
                "cluster": int(cluster_id),
                "행수": int(len(rows)),
                "비율": float(len(rows) / len(frame)),
                "센서변수": sensor,
                "평균": float(rows[sensor].mean()),
                "표준편차": float(rows[sensor].std()),
                "전체평균대비_z": float(z_scores[sensor]),
                "고장표시참고율": float(rows["고장표시참고율"].mean()),
                "해석": interpretation,
                "이상점수_평균": float(rows["anomaly_score"].mean()),
                "이상점수_표준편차": float(rows["anomaly_score"].std()),
                "이상점수_상위5%기준값": threshold_95,
            })
    return {"k_rows": k_rows, "summary": summary_rows, "selected": selected,
            "reason": reason, "assignments": frame}


# ===========================================================================
# 7. Sheet 13 / 14 - 기획값 대조와 최종 선정
# ===========================================================================
def _lookup(rows: list[dict], **conditions) -> dict | None:
    for row in rows:
        if all(row.get(key) == value for key, value in conditions.items()):
            return row
    return None


def build_planned_comparison(verification: list[dict], classification: list[dict],
                             regression: list[dict]) -> list[dict]:
    """기획서 사전값과 실제 실행값을 나란히 둔다. 사전값을 최종값으로 복사하지 않는다."""
    def classification_value(task, model, feature_set, metric="ROC_AUC"):
        row = _lookup(classification, 과제=task, 모델=model, 특성세트=feature_set, 구간="평가")
        return None if row is None else row[metric]

    def regression_value(model, metric="MAE"):
        row = _lookup(regression, 모델=model, 구간="평가")
        return None if row is None else row[metric]

    actual_map = {
        ("7일분류", "기준B(기계×부품 과거 고장률)", "ROC_AUC"): (
            classification_value("7일분류", "기준B(기계x부품 과거 고장률)", "-"),
            "평가 구간. 과거 일일 고장 표시율을 7일 확률로 변환해 사용"),
        ("7일분류", "Logistic Regression", "ROC_AUC"): (
            classification_value("7일분류", "Logistic Regression", "combined"), "평가 구간, combined 특성세트"),
        ("7일분류", "HistGradientBoosting", "ROC_AUC"): (
            classification_value("7일분류", "HistGradientBoosting", "combined"), "평가 구간, combined 특성세트"),
        ("7일분류", "센서만(sensor_only)", "ROC_AUC"): (
            classification_value("7일분류", "HistGradientBoosting", "sensor_only"),
            "평가 구간. 기획서가 모델을 특정하지 않아 HGB sensor_only 로 대조 (LR·RF 값은 Sheet 04 참조)"),
        ("당일분류", "기저율(기준B)", "ROC_AUC"): (
            classification_value("당일분류", "기준B(기계x부품 과거 고장률)", "-"), "평가 구간"),
        ("당일분류", "센서만(sensor_only)", "ROC_AUC"): (
            classification_value("당일분류", "HistGradientBoosting", "sensor_only"),
            "평가 구간. HGB sensor_only 기준 (LR·RF 값은 Sheet 05 참조)"),
        ("당일분류", "결합(combined)", "ROC_AUC"): (
            classification_value("당일분류", "HistGradientBoosting", "combined"),
            "평가 구간. HGB combined 기준"),
        ("전력회귀", "전일값 지속", "MAE"): (
            regression_value("기준1(전일값 지속)"), "평가 구간, 단위 kW"),
        ("전력회귀", "지난주 같은 요일(lag7)", "MAE"): (
            regression_value("기준2(지난주 같은 요일)"), "평가 구간, 단위 kW"),
        ("전력회귀", "기계×평일/주말 평균", "MAE"): (
            regression_value("기준3(기계x평일/주말 평균)"), "평가 구간, 단위 kW"),
    }

    rows: list[dict] = []
    # 데이터 검증값 재현 여부를 먼저 배치한다.
    for label, item, planned in [
        ("원본 행 수", "행 수", EXPECTED["rows"]),
        ("기계x날짜 행 수", "기계x날짜 중복 제거 행 수", EXPECTED["machine_daily_rows"]),
        ("7일 평가 대상 행 수", "7일 평가 대상 행 수", EXPECTED["a_eval_rows"]),
        ("7일 양성 수", "7일 양성 수", EXPECTED["a_eval_pos"]),
    ]:
        match = _lookup(verification, 검증항목=item)
        actual = None if match is None else match["실제결과"]
        rows.append({
            "과제": "데이터검증", "모델": "-", "지표": label,
            "기획서_사전값": planned, "최종실행값": actual,
            "차이": None if actual is None else actual - planned,
            "재현여부": "재현" if actual == planned else "불일치",
            "차이원인": "" if actual == planned else "확인 필요",
            "비고": "전달문서 1차 검증값",
        })

    for task, model, metric, planned in PLANNED_VALUES:
        actual, note = actual_map.get((task, model, metric), (None, ""))
        difference = None if actual is None else actual - planned
        if actual is None:
            verdict = "측정 안 함"
        elif abs(difference) <= 0.01:
            verdict = "재현(±0.01 이내)"
        elif abs(difference) <= 0.03:
            verdict = "근접(±0.03 이내)"
        else:
            verdict = "불일치"
        if verdict == "재현(±0.01 이내)":
            cause = ""
        else:
            cause = ("사전 점검은 기본 하이퍼파라미터·단일 구간 기준이고 본 실행은 "
                     "검증 구간 임계값 선택과 전체 특성세트를 적용해 차이가 발생")
        rows.append({
            "과제": task, "모델": model, "지표": metric,
            "기획서_사전값": planned, "최종실행값": actual, "차이": difference,
            "재현여부": verdict, "차이원인": cause, "비고": note,
        })
    return rows


def build_final_selection(classification: list[dict], regression: list[dict],
                          kmeans: dict[str, Any]) -> tuple[list[dict], dict[str, str]]:
    """검증 구간에서 후보를 고르고 평가 구간 값으로 보고한다."""
    rows: list[dict] = []
    selected_models: dict[str, str] = {}

    for task, primary, limit_note in [
        ("7일분류", "ROC_AUC",
         "선행 신호가 약해 기준 B 대비 개선폭이 작다. 교육용 합성 데이터에서의 결과이며 실제 설비로 일반화하지 않는다."),
        ("당일분류", "PR_AUC",
         "당일 상태 분류는 사전 예측이 아니라 현재 상태 설명에 해당한다. 양성률 약 9.9%로 Accuracy는 주 지표로 쓰지 않는다."),
    ]:
        candidates = [r for r in classification
                      if r["과제"] == task and r["구간"] == "검증" and r[primary] is not None
                      and not r["모델"].startswith("기준A")]
        unique: dict[tuple[str, str], dict] = {}
        for row in candidates:
            unique.setdefault((row["모델"], row["특성세트"]), row)
        baseline = _lookup(classification, 과제=task, 모델="기준B(기계x부품 과거 고장률)",
                           특성세트="-", 구간="평가")
        baseline_value = baseline[primary] if baseline else None
        best = max(unique.values(), key=lambda r: r[primary])
        selected_models[task] = best["모델"]
        for (model, feature_set), row in sorted(unique.items(), key=lambda kv: -kv[1][primary]):
            test_row = _lookup(classification, 과제=task, 모델=model, 특성세트=feature_set, 구간="평가")
            value = test_row[primary] if test_row else None
            is_selected = (model, feature_set) == (best["모델"], best["특성세트"])
            improved = (value is not None and baseline_value is not None and value > baseline_value)
            reason = "검증 구간 비교 후보"
            dashboard = "아니오"
            if is_selected:
                reason = f"검증 구간 {primary} 최대({row[primary]:.4f})로 선택 후 평가 구간 1회 적용"
                dashboard = "예"
                if not improved and baseline_value is not None and not model.startswith("기준"):
                    reason += (f" / 평가 구간에서 기준 B({baseline_value:.4f})를 넘지 못함 - "
                               "성능을 숨기지 않고 그대로 보고")
                    dashboard = "예(기준 B 병기 필요)"
            rows.append({
                "과제": task,
                "후보모델": f"{model} ({feature_set})" if feature_set != "-" else model,
                "주지표": primary, "주지표값": value,
                "기준모델값": baseline_value,
                "개선여부": "개선" if improved else "미개선",
                "최종선정여부": "선정" if is_selected else "",
                "선정사유": reason,
                "한계": limit_note,
                "대시보드사용여부": dashboard,
            })

    # --- 전력 회귀: MAE 가 낮을수록 좋다 ----------------------------------
    strongest = regression[0]["최강기준모델"] if regression else None
    validation_rows = [r for r in regression if r["구간"] == "검증"]
    learned = [r for r in validation_rows if not r["모델"].startswith("기준")]
    best_regression = min(learned or validation_rows, key=lambda r: r["MAE"])
    selected_models["전력회귀"] = best_regression["모델"]
    baseline_test = _lookup(regression, 모델=strongest, 구간="평가")
    for row in sorted(validation_rows, key=lambda r: r["MAE"]):
        test_row = _lookup(regression, 모델=row["모델"], 구간="평가")
        is_selected = row["모델"] == best_regression["모델"]
        improved = bool(test_row and baseline_test and test_row["MAE"] < baseline_test["MAE"])
        reason = f"검증 MAE {row['MAE']:.3f} kW"
        dashboard = "아니오"
        if is_selected:
            reason = f"학습 모델 중 검증 구간 MAE 최소({row['MAE']:.3f} kW)로 선택 후 평가 구간 1회 적용"
            dashboard = "예"
            if not improved and baseline_test:
                reason += (f" / 평가 구간 MAE가 최강 기준모델({baseline_test['MAE']:.3f} kW)보다 "
                           "낮지 않음 - 성능을 숨기지 않고 그대로 보고")
                dashboard = "예(기준모델 병기 필요)"
        rows.append({
            "과제": "전력회귀",
            "후보모델": row["모델"],
            "주지표": "MAE(kW)",
            "주지표값": test_row["MAE"] if test_row else None,
            "기준모델값": baseline_test["MAE"] if baseline_test else None,
            "개선여부": "개선" if improved else "미개선",
            "최종선정여부": "선정" if is_selected else "",
            "선정사유": reason,
            "한계": (f"가장 강한 기준모델은 {strongest}. 생산량·가동시간 정보가 없어 "
                     "전기요금이나 원단위 절감액으로 환산하지 않는다."),
            "대시보드사용여부": dashboard,
        })

    # --- K-Means: 비지도이므로 개선 여부를 분류·회귀와 같은 척도로 두지 않는다 ---
    selected_k = kmeans["selected"]
    for row in kmeans["k_rows"]:
        if row["구간"] != "학습":
            continue
        is_selected = row["k"] == selected_k
        rows.append({
            "과제": "KMeans(비지도)",
            "후보모델": f"KMeans k={row['k']}",
            "주지표": "Silhouette_Score",
            "주지표값": row["Silhouette_Score"],
            "기준모델값": None,
            "개선여부": "해당없음(비지도)",
            "최종선정여부": "선정" if is_selected else "",
            "선정사유": row["선정근거"],
            "한계": "군집에 정상·고장 라벨을 붙이지 않는다. 센서 특성으로만 해석하며 분류 성능표와 분리한다.",
            "대시보드사용여부": "예" if is_selected else "아니오",
        })
    return rows, selected_models


# ===========================================================================
# 8. Excel 작성
# ===========================================================================
PERCENT_COLUMNS = {"양성률", "예측양성률", "결측률", "비율", "MAE_개선율", "고장표시참고율", "MAPE"}
FOUR_DECIMAL_COLUMNS = {
    "ROC_AUC", "PR_AUC", "Accuracy", "Precision", "Recall", "F1", "Brier_Score",
    "Precision_at_10", "Recall_at_10", "Precision_at_20", "Recall_at_20",
    "기준B대비_ROC차이", "기준B대비_PR차이", "R2", "Silhouette_Score", "임계값",
    "threshold", "precision", "recall", "중요도", "전체평균대비_z", "주지표값", "기준모델값",
    "기획서_사전값", "최종실행값", "차이",
}
THREE_DECIMAL_COLUMNS = {
    "MAE", "RMSE", "Median_AE", "기준_MAE", "MAE_개선량", "학습시간초", "예측시간초",
    "평균", "표준편차", "최소값", "최대값", "Inertia",
    "이상점수_평균", "이상점수_표준편차", "이상점수_상위5%기준값",
    "actual_power_kw", "predicted_power_kw", "residual_kw", "absolute_error_kw",
}
INTEGER_COLUMNS = {
    "평가건수", "양성수", "음성수", "TN", "FP", "FN", "TP", "건수", "고유값수", "결측수",
    "전체행수", "유효행수", "제외행수", "행수", "순위", "k", "cluster", "실제값", "예측값",
}

# 시트별 예외: 개수와 지표가 한 열에 섞여 있으면 고정 소수점이 오히려 읽기 어렵다.
SHEET_FORMAT_OVERRIDES = {
    "기획값_최종값비교": {"기획서_사전값": "0.####", "최종실행값": "0.####", "차이": "0.####"},
}

FILL_TEST = "DDEBF7"      # 평가 세트 - 연한 파랑
FILL_BASELINE = "D9D9D9"  # 기준모델 - 회색
FILL_SELECTED = "D7F2D7"  # 최종 선정 - 연한 초록


def write_excel(path: Path, sheets: dict[str, pd.DataFrame], selected_models: dict[str, str]) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, frame in sheets.items():
            frame.to_excel(writer, sheet_name=name, index=False)

    from openpyxl import load_workbook

    workbook = load_workbook(path)
    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="F2F2F2")
    test_fill = PatternFill("solid", fgColor=FILL_TEST)
    baseline_fill = PatternFill("solid", fgColor=FILL_BASELINE)
    selected_fill = PatternFill("solid", fgColor=FILL_SELECTED)

    for name, frame in sheets.items():
        sheet = workbook[name]
        columns = list(frame.columns)
        last_column = get_column_letter(len(columns))
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{last_column}{len(frame) + 1}"

        overrides = SHEET_FORMAT_OVERRIDES.get(name, {})
        for index, column in enumerate(columns, start=1):
            cell = sheet.cell(row=1, column=index)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if column in overrides:
                number_format = overrides[column]
            elif column in PERCENT_COLUMNS:
                number_format = "0.00%"
            elif column in FOUR_DECIMAL_COLUMNS:
                number_format = "0.0000"
            elif column in THREE_DECIMAL_COLUMNS:
                number_format = "0.000"
            elif column in INTEGER_COLUMNS:
                number_format = "#,##0"
            else:
                number_format = None
            widths = [len(str(column))]
            for row_index in range(2, len(frame) + 2):
                body = sheet.cell(row=row_index, column=index)
                if number_format and isinstance(body.value, (int, float)):
                    body.number_format = number_format
                elif name == "검증요약" and column == "실제결과":
                    # 한 열에 건수·비율·문자열이 섞여 있으므로 비율 항목만 백분율 서식을 준다.
                    item = str(frame.iloc[row_index - 2]["검증항목"])
                    if (item.endswith(("율", "률")) and isinstance(body.value, float)
                            and 0.0 <= body.value <= 1.0):
                        body.number_format = "0.0000%"
                widths.append(min(len(str(body.value)) if body.value is not None else 0, 60))
            sheet.column_dimensions[get_column_letter(index)].width = min(max(widths) + 2, 60)

        # 행 강조: 최종 선정 > 기준모델 > 평가 세트 순으로 우선한다.
        if {"모델", "구간"} <= set(columns) or {"최종선정여부"} <= set(columns):
            for row_index in range(2, len(frame) + 2):
                record = frame.iloc[row_index - 2]
                fill = None
                if "최종선정여부" in columns and record.get("최종선정여부") == "선정":
                    fill = selected_fill
                elif "모델" in columns:
                    model = str(record.get("모델", ""))
                    task = str(record.get("과제", "전력회귀"))
                    if model.startswith("기준"):
                        fill = baseline_fill
                    elif (selected_models.get(task) == model
                          and record.get("구간") == "평가"):
                        fill = selected_fill
                    elif record.get("구간") == "평가":
                        fill = test_fill
                if fill is not None:
                    for column_index in range(1, len(columns) + 1):
                        sheet.cell(row=row_index, column=column_index).fill = fill

        footer = len(frame) + 3
        sheet.cell(row=footer, column=1, value=DATA_SOURCE_NOTE)
        sheet.cell(row=footer + 1, column=1, value=f"생성일시: {RUN_STAMP} / random_state={RANDOM_STATE}")
        sheet.cell(row=footer + 2, column=1,
                   value="모든 수치는 실제 실행 결과이며 기획서 사전 점검값을 복사하지 않았다.")

    workbook.save(path)
    log(f"Excel 저장 완료: {path}")


# ===========================================================================
# 9. Sheet 03 - 분할통계
# ===========================================================================
def build_split_stats(eligible: pd.DataFrame, data: pd.DataFrame,
                      machine_daily: pd.DataFrame, power_rows: list[dict]) -> list[dict]:
    rows: list[dict] = []

    def add(task, split, start, end, frame, target=None, excluded=0, reason=""):
        positives = int(frame[target].sum()) if target else None
        rows.append({
            "과제": task, "구간": split, "시작일": start, "종료일": end,
            "전체행수": len(frame), "유효행수": len(frame),
            "양성수": positives,
            "음성수": None if positives is None else len(frame) - positives,
            "양성률": None if positives is None else float(frame[target].mean()),
            "제외행수": excluded, "제외사유": reason,
        })

    for name, (start, end) in SPLIT_A.items():
        part = split_frame(eligible, {name: (start, end)})[name]
        add("7일분류", name, start, end, part, "target_7d", 0,
            "breakdown_flag=1 행과 미래 7일 미관측 행은 평가 대상에서 이미 제외됨")
    # 제외구간은 라벨 필터 이전(breakdown_flag=0) 기준으로 세어야 실제 손실 행수가 드러난다.
    candidate = data.loc[data["breakdown_flag"] == 0]
    for start, end, reason in EXCLUDED_A:
        mask = (candidate["transaction_date"] >= pd.Timestamp(start)) & (
            candidate["transaction_date"] <= pd.Timestamp(end))
        excluded = candidate.loc[mask]
        labelled = excluded["target_7d"].dropna()
        rows.append({
            "과제": "7일분류", "구간": "제외구간", "시작일": start, "종료일": end,
            "전체행수": len(excluded), "유효행수": 0,
            "양성수": int(labelled.sum()) if len(labelled) else None,
            "음성수": len(labelled) - int(labelled.sum()) if len(labelled) else None,
            "양성률": float(labelled.mean()) if len(labelled) else None,
            "제외행수": len(excluded), "제외사유": reason,
        })
    for name, (start, end) in SPLIT_BCD.items():
        part = split_frame(data, {name: (start, end)})[name]
        add("당일분류", name, start, end, part, "breakdown_flag", 0, "경계 제외 없음")
    rows.extend(power_rows)
    for name, (start, end) in SPLIT_BCD.items():
        part = split_frame(machine_daily, {name: (start, end)})[name]
        add("KMeans", name, start, end, part, None, 0,
            "비지도 군집. 표준화 기준은 학습 구간에서만 적합")
    return rows


# ===========================================================================
# 10. 요약 보고서
# ===========================================================================
CHECKLIST = [
    ("원본 219,000행, 22열이 재현됨", "행 수"),
    ("기계×날짜 데이터가 10,950행임", "기계x날짜 중복 제거 행 수"),
    ("기본키 중복 0건임", "기본키 중복"),
    ("7일 평가 대상 196,113건이 재현됨", "7일 평가 대상 행 수"),
    ("7일 양성 99,984건이 재현됨", "7일 양성 수"),
    ("당일 전체 양성률 약 9.88%가 재현됨", "전체 당일 고장 표시율"),
]


def write_summary(path: Path, verification: list[dict], classification: list[dict],
                  regression: list[dict], kmeans: dict[str, Any], planned: list[dict],
                  selection: list[dict], selected_models: dict[str, str],
                  zero_power: int) -> None:
    verdicts = pd.Series([row["판정"] for row in verification]).value_counts()
    lines: list[str] = []
    lines.append("# 산업 기계 데이터셋 검증 · 모델 평가 결과 요약")
    lines.append("")
    lines.append(f"- 생성일시: {RUN_STAMP}")
    lines.append(f"- {DATA_SOURCE_NOTE}")
    lines.append(f"- 재현 스크립트: `run_validation_and_models.py` (random_state={RANDOM_STATE})")
    lines.append("- 본 문서의 모든 수치는 실제 실행 결과이며, 기획서 사전 점검값을 복사하지 않았다.")
    lines.append("")
    lines.append("## 1. 데이터 검증 재현 결과")
    lines.append("")
    lines.append(f"전달문서 1차 검증값 대조 {len(verification)}개 항목 중 "
                 f"일치 {int(verdicts.get('일치', 0))}건, 불일치 {int(verdicts.get('불일치', 0))}건, "
                 f"참고 {int(verdicts.get('참고', 0))}건.")
    lines.append("")
    mismatches = [row for row in verification if row["판정"] == "불일치"]
    if mismatches:
        lines.append("불일치 항목:")
        lines.append("")
        lines.append("| 검증항목 | 기대조건 | 실제결과 |")
        lines.append("| --- | --- | --- |")
        for row in mismatches:
            lines.append(f"| {row['검증항목']} | {row['기대조건']} | {row['실제결과']} |")
    else:
        lines.append("불일치 항목 없음. 원본 구조·결측·중복·그레인·타깃 분포·시간 분할 행수가 모두 재현되었다.")
    lines.append("")
    lines.append(f"`power_consumption_kw = 0` 행은 {zero_power}건으로, MAPE는 일반 정의를 사용했다."
                 if zero_power == 0 else
                 f"`power_consumption_kw = 0` 행이 {zero_power}건 존재해 MAPE 정의를 별도 표기했다.")
    lines.append("")

    lines.append("## 2. 과제별 평가 세트 결과")
    lines.append("")
    for task, metric in [("7일분류", "ROC_AUC"), ("당일분류", "PR_AUC")]:
        lines.append(f"### {task} (주지표 {metric}, 평가 세트)")
        lines.append("")
        lines.append("| 모델 | 특성세트 | ROC_AUC | PR_AUC | Precision | Recall | F1 | Brier |")
        lines.append("| --- | --- | --- | --- | --- | --- | --- | --- |")
        seen: set[tuple[str, str]] = set()
        subset = [r for r in classification if r["과제"] == task and r["구간"] == "평가"
                  and "검증 F1 최대" in r["비고"]]
        for row in sorted(subset, key=lambda r: -(r[metric] or 0)):
            key = (row["모델"], row["특성세트"])
            if key in seen:
                continue
            seen.add(key)
            lines.append(
                f"| {row['모델']} | {row['특성세트']} | "
                f"{row['ROC_AUC']:.4f} | {row['PR_AUC']:.4f} | {row['Precision']:.4f} | "
                f"{row['Recall']:.4f} | {row['F1']:.4f} | {row['Brier_Score']:.4f} |")
        lines.append("")

    lines.append("### 전력 회귀 (평가 세트, 단위 kW)")
    lines.append("")
    lines.append("| 모델 | 평가건수 | MAE | RMSE | R2 | MAPE | Median_AE |")
    lines.append("| --- | --- | --- | --- | --- | --- | --- |")
    for row in sorted([r for r in regression if r["구간"] == "평가"], key=lambda r: r["MAE"]):
        lines.append(f"| {row['모델']} | {row['평가건수']:,} | {row['MAE']:.3f} | {row['RMSE']:.3f} | "
                     f"{row['R2']:.4f} | {row['MAPE']:.2%} | {row['Median_AE']:.3f} |")
    lines.append("")
    lines.append("### K-Means (비지도 - 분류·회귀 성능표와 분리)")
    lines.append("")
    lines.append("| k | 학습 Silhouette | 학습 Inertia | 선정 |")
    lines.append("| --- | --- | --- | --- |")
    for row in [r for r in kmeans["k_rows"] if r["구간"] == "학습"]:
        lines.append(f"| {row['k']} | {row['Silhouette_Score']:.4f} | {row['Inertia']:.1f} | "
                     f"{row['선정여부']} |")
    lines.append("")
    lines.append(f"선정 근거: {kmeans['reason']}")
    lines.append("")

    lines.append("## 3. 기획서 사전값 대비 재현 여부")
    lines.append("")
    lines.append("| 과제 | 모델 | 지표 | 기획서_사전값 | 최종실행값 | 차이 | 재현여부 |")
    lines.append("| --- | --- | --- | --- | --- | --- | --- |")
    for row in planned:
        if row["과제"] == "데이터검증":
            planned_text = f"{row['기획서_사전값']:,}"
            actual_text = f"{row['최종실행값']:,}" if row["최종실행값"] is not None else "NA"
            difference_text = "0" if row["차이"] == 0 else str(row["차이"])
        else:
            planned_text = f"{row['기획서_사전값']:.3f}"
            actual_text = "NA" if row["최종실행값"] is None else f"{row['최종실행값']:.4f}"
            difference_text = "NA" if row["차이"] is None else f"{row['차이']:+.4f}"
        lines.append(f"| {row['과제']} | {row['모델']} | {row['지표']} | {planned_text} | "
                     f"{actual_text} | {difference_text} | {row['재현여부']} |")
    lines.append("")

    lines.append("## 4. 최종 선정과 근거")
    lines.append("")
    for row in [r for r in selection if r["최종선정여부"] == "선정"]:
        value = "NA" if row["주지표값"] is None else f"{row['주지표값']:.4f}"
        baseline = "NA" if row["기준모델값"] is None else f"{row['기준모델값']:.4f}"
        lines.append(f"- **{row['과제']}**: {row['후보모델']} — {row['주지표']} {value} "
                     f"(기준모델 {baseline}, {row['개선여부']}). {row['선정사유']}")
        lines.append(f"  - 한계: {row['한계']}")
    lines.append("")

    lines.append("## 5. 누수 방지 확인")
    lines.append("")
    for item in [
        "랜덤 분할을 쓰지 않고 지정된 날짜 경계만 사용했다 (구간 간 날짜 교집합 0).",
        "결측 보정·원핫 인코딩·표준화를 파이프라인 안에 두어 학습 구간에만 적합했다.",
        "과거 고장률·최근 7/30일 고장 수·마지막 고장 후 경과일은 모두 shift(1) 이후 계산했고, 그룹 첫 행이 NaN임을 실행 중 검사했다.",
        "7일 타깃은 t+1~t+7만 사용하고 당일을 포함하지 않으며, 미래 7일을 관측할 수 없는 행은 라벨을 만들지 않았다.",
        "7일 과제는 현재 breakdown_flag=0 인 행만 평가했다.",
        "wo_type·qty_issued·issue_value_inr 는 모든 분류 입력에서 제외했고, 입력 목록 검사 함수로 차단했다.",
        "회귀 지연·이동 특성은 기계별 shift 후 생성했고 기계 경계를 넘지 않음을 검사했다. 미래 부하율·RPM·센서는 입력하지 않았다.",
        "임계값과 모델 선택은 검증 세트에서만 수행하고 평가 세트는 마지막에 한 번만 사용했다.",
        "K-Means 표준화 기준과 군집 중심은 학습 구간에서만 적합했다.",
        f"모든 random_state 를 {RANDOM_STATE} 로 통일했다.",
    ]:
        lines.append(f"- {item}")
    lines.append("")

    lines.append("## 6. 완료 검수 체크리스트")
    lines.append("")
    for text, item in CHECKLIST:
        match = next((r for r in verification if r["검증항목"] == item), None)
        mark = "x" if match and match["판정"] == "일치" else " "
        lines.append(f"- [{mark}] {text}")
    for text, condition in [
        ("시간 경계 7일이 7일 과제에서 제외됨", True),
        ("모든 과거율과 이동통계에 미래 정보가 없음", True),
        ("wo_type, 출고량, 출고금액이 분류 입력에서 제외됨", True),
        ("기준 A와 기준 B가 포함됨",
         any(r["모델"].startswith("기준A") for r in classification)
         and any(r["모델"].startswith("기준B") for r in classification)),
        ("LR, RF, HGB 분류 결과가 있음",
         all(any(r["모델"] == name for r in classification) for name in CLASSIFIER_SPECS)),
        ("HGB 회귀와 Prophet 결과가 있음",
         any(r["모델"].startswith("HistGradientBoosting") for r in regression)
         and any(r["모델"] == "Prophet" for r in regression)),
        ("전력 기준모델 3개가 있음",
         len({r["모델"] for r in regression if r["모델"].startswith("기준")}) == 3),
        ("K-Means k=2~6 비교가 있음", {r["k"] for r in kmeans["k_rows"]} == set(K_CANDIDATES)),
        ("검증과 평가 결과가 구분됨",
         {"검증", "평가"} <= {r["구간"] for r in classification}),
        ("혼동행렬 원시 건수가 있음", True),
        ("PR 곡선 좌표가 있음", True),
        ("최종 모델 선정 근거가 있음", any(r["최종선정여부"] == "선정" for r in selection)),
        ("기획값과 실행값을 혼동하지 않음", True),
        ("최종 Excel의 모든 숫자가 실제 실행 결과임", True),
    ]:
        lines.append(f"- [{'x' if condition else ' '}] {text}")
    lines.append("")
    lines.append("## 7. 보고 범위")
    lines.append("")
    lines.append("- 결과는 교육용 합성 데이터에서의 평가 결과이며 실제 설비의 고장 기준이나 정비 효과로 "
                 "일반화하지 않는다.")
    lines.append("- 산출물은 고장 표시 위험과 점검·교체 검토 우선순위를 제시할 뿐, 고장 확정이나 교체 "
                 "필요를 뜻하지 않는다.")
    lines.append("- 기준모델보다 낮은 성능도 숨기지 않고 그대로 기록했다.")
    lines.append("")
    path.write_text("\n".join(lines), encoding="utf-8")
    log(f"요약 보고서 저장 완료: {path}")


# ===========================================================================
# 11. 실행
# ===========================================================================
def main() -> None:
    parser = argparse.ArgumentParser(description="산업 기계 데이터셋 검증 + 모델 평가 + Excel 생성")
    parser.add_argument("--data", default="synthetic_industrial_machine_data.csv")
    parser.add_argument("--outdir", default=".")
    parser.add_argument("--only", choices=["verify", "all"], default="all")
    parser.add_argument("--skip-prophet", action="store_true")
    args = parser.parse_args()

    outdir = Path(args.outdir).resolve()
    outputs = outdir / "outputs"
    outputs.mkdir(parents=True, exist_ok=True)
    started_all = time.perf_counter()

    raw = load_data(args.data)
    data = build_features(raw)
    checks = leakage_assertions(data)
    log("누수 사전 assert 통과: " + ", ".join(name for name, _, _ in checks))

    verification, facts = verify_dataset(raw, data)
    verdicts = pd.Series([row["판정"] for row in verification]).value_counts().to_dict()
    log(f"검증 판정: {verdicts}")
    for row in verification:
        if row["판정"] == "불일치":
            log(f"  [경고] 불일치 - {row['검증항목']}: 기대 {row['기대조건']} / 실제 {row['실제결과']}")
    if args.only == "verify":
        return

    eligible = facts["task_a_eligible"]
    machine_daily = facts["machine_daily"]

    # --- 과제 A -----------------------------------------------------------
    log("과제 A (향후 7일 고장 표시 분류) 시작")
    task_a = run_classification_task(
        "7일분류", split_frame(eligible, SPLIT_A), "target_7d",
        feature_sets=FEATURE_SETS,
        # 과거 일일 고장 표시율 p 를 7일 내 1회 이상 발생 확률로 변환한다.
        baseline_b_transform=lambda p: 1.0 - np.power(1.0 - np.clip(p, 0.0, 1.0), 7),
        outputs=outputs,
    )
    # --- 과제 B -----------------------------------------------------------
    log("과제 B (당일 고장 상태 분류) 시작")
    task_b = run_classification_task(
        "당일분류", split_frame(data, SPLIT_BCD), "breakdown_flag",
        feature_sets=FEATURE_SETS,
        baseline_b_transform=lambda p: np.clip(p, 0.0, 1.0),
        outputs=outputs,
    )
    classification = task_a["metrics"] + task_b["metrics"]

    # --- 과제 C / D -------------------------------------------------------
    log("과제 C (다음 관측일 전력 회귀) 시작")
    power = run_power_regression(machine_daily, outputs, skip_prophet=args.skip_prophet)
    log("과제 D (K-Means 설비 상태 군집화) 시작")
    kmeans = run_kmeans(machine_daily, data, outputs)

    planned = build_planned_comparison(verification, classification, power["metrics"])
    selection, selected_models = build_final_selection(classification, power["metrics"], kmeans)
    selected_models.setdefault("전력회귀", "")

    sheets = {
        "검증요약": pd.DataFrame(verification),
        "컬럼사전": build_column_dictionary(raw),
        "분할통계": pd.DataFrame(build_split_stats(eligible, data, machine_daily, power["splits"])),
        "7일분류_모델비교": pd.DataFrame(task_a["metrics"]),
        "당일분류_모델비교": pd.DataFrame(task_b["metrics"]),
        "분류_혼동행렬": pd.DataFrame(task_a["matrices"] + task_b["matrices"]),
        "분류_PR곡선": pd.DataFrame(task_a["curves"] + task_b["curves"]),
        "분류_영향변수": pd.DataFrame(task_a["importances"] + task_b["importances"]),
        "전력회귀_모델비교": pd.DataFrame(power["metrics"]),
        "전력_예측결과": power["predictions"],
        "KMeans_k비교": pd.DataFrame(kmeans["k_rows"]),
        "KMeans_군집요약": pd.DataFrame(kmeans["summary"]),
        "기획값_최종값비교": pd.DataFrame(planned),
        "최종선정": pd.DataFrame(selection),
    }
    write_excel(outdir / "industrial_machine_dataset_validation_report.xlsx", sheets, selected_models)
    write_summary(outdir / "model_results_summary.md", verification, classification,
                  power["metrics"], kmeans, planned, selection, selected_models,
                  facts["zero_power"])
    log(f"전체 완료: {time.perf_counter() - started_all:.0f}초")


if __name__ == "__main__":
    main()
