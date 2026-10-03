"""보고서 내보내기 — 독자별 섹션 골격의 PDF·Excel."""

import io
from datetime import datetime

from ..theme import C
from .base import (
    _period_index, DEFAULT_AUDIENCE, DEFAULT_FILTERS, DEFAULT_SEG, NO_DATA_MARK, PERIOD_PRESETS,
    REPORT_AUDIENCES, SEG_GROUPS,
)


# ============================================================
# 보고서 내보내기 (PDF / Excel)
#
# 데이터가 아직 붙지 않았으므로 산출물에 가짜 숫자를 넣지 않는다. 대신
# "이 독자에게는 어떤 섹션·어떤 열이 들어가는지"가 확정된 보고서 골격을
# 내보낸다. 값 칸은 전부 NO_DATA_MARK다. 데이터가 붙으면 _section_rows()만
# 실제 쿼리 결과로 바꾸면 되고, 레이아웃·섹션 구성은 그대로 쓴다.
# ============================================================

CAVEAT = ("합성 데이터 · 교육용 — 이 보고서의 모든 수치는 실제 설비 이력이 아니다. "
          "운영 판단의 단독 근거로 쓰지 않는다.")

# ------------------------------------------------------------
# PDF 한글 폰트
#
# reportlab 내장 CID 폰트(HYSMyeongJo-Medium)는 쓰지 않는다. 폰트를 PDF에
# 임베드하지 않고 뷰어의 Adobe 한국어 폰트팩에 의존하기 때문에, 그 팩이 없는
# 뷰어(크롬 내장 뷰어, poppler, 미리보기 등)에서는 글자가 통째로 사라진다.
# 실제로 그 증상을 확인했다. 그래서 시스템에 있는 한글 TrueType을 찾아
# 임베드한다. TrueType만 된다 — Noto Sans CJK 같은 OTF/CFF 계열은 reportlab이
# 거부한다("postscript outlines are not supported").
#
# 하나도 못 찾으면 PDF를 만들지 않고 실패로 알린다. 글자 없는 PDF를
# 성공인 척 내보내지 않는다.
# ------------------------------------------------------------
KOREAN_TTF_CANDIDATES = [
    # Windows
    r"C:\Windows\Fonts\malgun.ttf",           # 맑은 고딕
    r"C:\Windows\Fonts\malgunsl.ttf",
    r"C:\Windows\Fonts\gulim.ttc",            # 굴림
    r"C:\Windows\Fonts\batang.ttc",           # 바탕
    # macOS
    "/System/Library/Fonts/Supplemental/AppleGothic.ttf",
    "/Library/Fonts/AppleGothic.ttf",
    "/System/Library/Fonts/AppleSDGothicNeo.ttc",
    # Linux (sudo apt-get install fonts-nanum)
    "/usr/share/fonts/truetype/nanum/NanumGothic.ttf",
    "/usr/share/fonts/truetype/nanum/NanumBarunGothic.ttf",
    "/usr/share/fonts/truetype/nanum/NanumMyeongjo.ttf",
    "/usr/share/fonts/nanum/NanumGothic.ttf",
    # 이 파일 옆에 폰트를 직접 두는 경우
    "NanumGothic.ttf", "malgun.ttf",
]
PDF_FONT_NAME = "KoreanBody"
_pdf_font_cache = {}


def resolve_pdf_font():
    """등록에 성공한 (폰트명, 경로)를 돌려준다. 못 찾으면 (None, None)."""
    if _pdf_font_cache:
        return _pdf_font_cache["name"], _pdf_font_cache["path"]
    import os
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    for path in KOREAN_TTF_CANDIDATES:
        if not os.path.exists(path):
            continue
        # .ttc는 한 파일에 여러 폰트가 들어 있어 인덱스를 훑어야 한다
        indices = (0, 1, 2, 3) if path.lower().endswith(".ttc") else (None,)
        for idx in indices:
            try:
                kwargs = {} if idx is None else {"subfontIndex": idx}
                pdfmetrics.registerFont(TTFont(PDF_FONT_NAME, path, **kwargs))
                # 한글 글리프가 실제로 있는지 확인 — 라틴 전용 폰트를 거른다
                if pdfmetrics.getFont(PDF_FONT_NAME).stringWidth("설비 모니터링", 10) <= 0:
                    raise ValueError("한글 글리프 없음")
            except Exception:
                continue
            _pdf_font_cache.update(name=PDF_FONT_NAME, path=path)
            return PDF_FONT_NAME, path
    return None, None


class PdfFontMissing(RuntimeError):
    pass


def _report_context(filters, seg_state, audience):
    filters = filters or DEFAULT_FILTERS
    seg_state = seg_state or DEFAULT_SEG
    aud = REPORT_AUDIENCES.get(audience, REPORT_AUDIENCES[DEFAULT_AUDIENCE])
    return {
        "aud": aud,
        "generated": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "meta": [
            ("보고서 대상", aud["label"]),
            ("보고 목적", aud["purpose"]),
            ("집계 단위", aud["grain"]),
            ("생성 시각", datetime.now().strftime("%Y-%m-%d %H:%M")),
            ("데이터 기준일", NO_DATA_MARK),
            ("공장", filters.get("plant") or "전체"),
            ("기계 종류", filters.get("machine_type") or "전체"),
            ("기계", filters.get("machine") or "전체"),
            ("기간", PERIOD_PRESETS[_period_index(filters)][0]),
            ("모델 과제", SEG_GROUPS["task"][seg_state.get("task", 0)]),
            ("위험 기준선", SEG_GROUPS["threshold"][seg_state.get("threshold", 0)]),
            ("제외 섹션", aud["excludes"]),
        ],
    }


def esc(s):
    """Paragraph는 미니 HTML을 파싱하므로 & < > 를 이스케이프한다."""
    from xml.sax.saxutils import escape
    return escape(str(s))


def _section_rows(cols):
    """데이터 연결 전까지 값 칸은 미연결 표시 1행. 연결 후 이 함수만 교체한다."""
    return [[NO_DATA_MARK] * len(cols)]


def _print_color(name):
    """PDF·Excel 산출물 색 — 화면과 같은 토큰(라이트 테마)에서 가져온다."""
    return C(name, "light")


def build_report_pdf(filters, seg_state, audience):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle)

    FONT, font_path = resolve_pdf_font()
    if not FONT:
        raise PdfFontMissing(
            "PDF에 임베드할 한글 TrueType 폰트를 찾지 못했다. "
            "Windows/macOS는 기본 폰트가 잡히고, Linux는 "
            "`sudo apt-get install fonts-nanum` 후 다시 시도하거나 "
            "NanumGothic.ttf를 이 파일 옆에 두면 된다. (Excel 내보내기는 영향 없음)")

    ctxd = _report_context(filters, seg_state, audience)
    aud = ctxd["aud"]

    h1 = ParagraphStyle("h1", fontName=FONT, fontSize=17, leading=22, spaceAfter=2, alignment=TA_LEFT)
    h2 = ParagraphStyle("h2", fontName=FONT, fontSize=11.5, leading=15, spaceBefore=9, spaceAfter=3)
    body = ParagraphStyle("body", fontName=FONT, fontSize=8.5, leading=12)
    small = ParagraphStyle("small", fontName=FONT, fontSize=7.5, leading=10.5,
                           textColor=colors.HexColor(_print_color("ink-muted")))

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=landscape(A4),
        leftMargin=14 * mm, rightMargin=14 * mm, topMargin=12 * mm, bottomMargin=12 * mm,
        title=f"설비 모니터링 보고서 — {aud['label']}", author="설비 모니터링 대시보드",
    )

    grid = TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), FONT),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor(_print_color("chart-axis"))),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(_print_color("surface-sunken"))),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("TEXTCOLOR", (0, 1), (-1, -1), colors.HexColor(_print_color("ink-muted"))),
    ])

    flow = [
        Paragraph("설비 모니터링 보고서", h1),
        Paragraph(esc(f"{aud['label']} · {aud['purpose']} · 생성 {ctxd['generated']}"), small),
        Spacer(1, 6),
        Paragraph(esc(CAVEAT), small),
        Spacer(1, 8),
        Paragraph("보고 조건", h2),
    ]

    meta_tbl = Table([[Paragraph(esc(k), body), Paragraph(esc(v), body)] for k, v in ctxd["meta"]],
                     colWidths=[34 * mm, 205 * mm])
    meta_tbl.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), FONT),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor(_print_color("chart-axis"))),
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor(_print_color("surface-sunken"))),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    flow.append(meta_tbl)

    flow.append(Paragraph(f"섹션 {len(aud['sections'])}개", h2))
    for i, (src, name, cols) in enumerate(aud["sections"], 1):
        flow.append(Paragraph(f"{esc(i)}. {esc(name)}　<font size=7.5>[{esc(src)}]</font>", h2))
        data = [[Paragraph(esc(c), body) for c in cols]] + [[Paragraph(esc(c), body) for c in r]
                                                        for r in _section_rows(cols)]
        avail = 239 * mm
        cw = [avail / len(cols)] * len(cols)
        t = Table(data, colWidths=cw, repeatRows=1)
        t.setStyle(grid)
        flow.append(t)

    flow.append(Spacer(1, 8))
    flow.append(Paragraph(
        f"값 칸이 '{NO_DATA_MARK}'인 것은 데이터 소스가 아직 연결되지 않았기 때문이다. "
        "섹션 구성과 열 정의는 확정본이므로, 데이터 연결 후 값만 채우면 된다.  "
        f"(임베드 폰트: {esc(font_path)})", small))

    doc.build(flow)
    return buf.getvalue()


def _safe_sheet(name, used):
    bad = set('[]:*?/\\')
    s = "".join(("·" if ch in bad else ch) for ch in name)[:31] or "sheet"
    base, n = s, 2
    while s in used:
        suffix = f"_{n}"
        s = base[:31 - len(suffix)] + suffix
        n += 1
    used.add(s)
    return s


def build_report_xlsx(filters, seg_state, audience):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    ctxd = _report_context(filters, seg_state, audience)
    aud = ctxd["aud"]

    def xl(name):  # openpyxl은 '#' 없는 RRGGBB를 받는다.
        return _print_color(name).lstrip("#").upper()

    thin = Side(style="thin", color=xl("chart-axis"))
    edge = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor=xl("surface-sunken"))
    bold = Font(bold=True, size=10)
    muted = Font(size=10, color=xl("ink-muted"))

    wb = Workbook()
    used = set()

    ws = wb.active
    ws.title = _safe_sheet("표지·보고조건", used)
    ws["A1"] = "설비 모니터링 보고서"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"{aud['label']} · {aud['purpose']}"
    ws["A2"].font = muted
    ws["A3"] = CAVEAT
    ws["A3"].font = muted
    r = 5
    ws.cell(r, 1, "항목").font = bold
    ws.cell(r, 2, "값").font = bold
    for c in (1, 2):
        ws.cell(r, c).fill = head_fill
        ws.cell(r, c).border = edge
    for k, v in ctxd["meta"]:
        r += 1
        ws.cell(r, 1, k).border = edge
        cell = ws.cell(r, 2, str(v))
        cell.border = edge
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    r += 2
    ws.cell(r, 1, "섹션 목록").font = bold
    for i, (src, name, cols) in enumerate(aud["sections"], 1):
        r += 1
        ws.cell(r, 1, f"{i}. {name}")
        ws.cell(r, 2, f"[{src}] · 열 {len(cols)}개")
        ws.cell(r, 2).font = muted
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 96
    ws.freeze_panes = "A6"

    for i, (src, name, cols) in enumerate(aud["sections"], 1):
        s = wb.create_sheet(_safe_sheet(f"{i}·{name}", used))
        s["A1"] = f"{name}  [{src}]"
        s["A1"].font = Font(bold=True, size=11)
        s["A2"] = f"{NO_DATA_MARK} — 열 정의는 확정본, 값은 데이터 연결 후 채운다."
        s["A2"].font = muted
        for c, label in enumerate(cols, 1):
            cell = s.cell(4, c, label)
            cell.font = bold
            cell.fill = head_fill
            cell.border = edge
            s.column_dimensions[get_column_letter(c)].width = max(14, min(30, len(label) + 6))
        for rr, rowvals in enumerate(_section_rows(cols), start=5):
            for c, v in enumerate(rowvals, 1):
                cell = s.cell(rr, c, v)
                cell.border = edge
                cell.font = muted
        s.freeze_panes = "A5"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
